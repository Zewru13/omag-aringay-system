"""Exports must never turn a name someone typed into a spreadsheet formula (which Excel would run when staff open the file)."""
import csv
import io
import zipfile
from datetime import date

import openpyxl
from django.contrib.auth import get_user_model
from django.http import HttpResponse
from django.test import SimpleTestCase
from rest_framework.test import APITestCase

from core.exports import SafeWorkbook, csv_safe_cell, safe_csv_writer
from farmers.models import Farmer
from production.models import CropMonitoringRecord, CropType, HarvestRecord

EVIL = '=HYPERLINK("http://evil.test/steal","Click to verify")'


def xml_of(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return ''.join(z.read(name).decode('utf-8', 'replace') for name in z.namelist() if name.endswith('.xml'))


class WorkbookTests(SimpleTestCase):
    def saved(self, *values):
        wb = SafeWorkbook()
        ws = wb.active
        for column, value in enumerate(values, start=1):
            ws.cell(row=1, column=column, value=value)
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    def test_text_that_looks_like_a_formula_is_stored_as_plain_text(self):
        data = self.saved(EVIL, '=1+1', '@SUM(A2)', '+cmd|calc')
        cells = openpyxl.load_workbook(io.BytesIO(data)).active[1]
        self.assertEqual([c.data_type for c in cells], ['s', 's', 's', 's'])
        self.assertEqual(cells[0].value, EVIL)                       # the person still sees exactly what was typed
        self.assertFalse('<f>' in xml_of(data), 'a formula was written')

    def test_numbers_dates_and_ordinary_text_are_untouched(self):
        data = self.saved(5, 2.5, date(2026, 1, 2), 'Maria Santos', None)
        cells = openpyxl.load_workbook(io.BytesIO(data)).active[1]
        self.assertEqual([cells[0].value, cells[1].value, cells[3].value], [5, 2.5, 'Maria Santos'])
        self.assertEqual(cells[0].data_type, 'n')


class CsvTests(SimpleTestCase):
    def test_dangerous_starts_are_defused_and_ordinary_values_are_not(self):
        for value in ('=1+1', '+cmd|calc', '-2+3', '@SUM(A1)', '\tcmd', '\rcmd', EVIL):
            self.assertEqual(csv_safe_cell(value), "'" + value, value)
        for value in ('Maria', '', '-5', '-5.25', '+639171234567', '12', 12, 3.5, None, date(2026, 1, 2)):
            self.assertEqual(csv_safe_cell(value), value, value)

    def test_the_writer_cleans_every_cell_of_a_row(self):
        response = HttpResponse()
        writer = safe_csv_writer(response)
        writer.writerow(['Name', EVIL, 7])
        writer.writerows([['ok', '=2+2']])
        rows = list(csv.reader(io.StringIO(response.content.decode())))
        self.assertEqual(rows, [['Name', "'" + EVIL, '7'], ['ok', "'=2+2"]])


class ReportsEndToEndTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = get_user_model().objects.create_user(username='xs-admin', password='pw', role='ADMIN', email='xs@x.test')
        cls.farmer = Farmer.objects.create(
            rsbsa_number='01-033-02-019-000777', first_name=EVIL, last_name='Santos', gender='Female',
            birthdate=date(1985, 3, 3), barangay='Alaska', civil_status='Single')
        crop = CropType.objects.order_by('id').first()
        cls.record = CropMonitoringRecord.objects.create(
            farmer=cls.farmer, farm_location_barangay='Alaska', land_status='Owner', crop_type=crop, season='Wet', year=2026,
            date_planted=date(2026, 5, 1), area_planted_ha='1.50', status='Harvested', actual_harvest_date=date(2026, 9, 1))
        HarvestRecord.objects.create(monitoring_record=cls.record, season='Wet', year=2026, crop_name=crop.name,
                                     area_harvested_hectares='1.50', volume_kg='3000.00')

    def setUp(self):
        self.client.force_authenticate(self.admin)

    def test_the_excel_report_keeps_a_formula_like_name_as_text(self):
        response = self.client.get('/api/reports/production/harvested-crop/')
        self.assertEqual(response.status_code, 200)
        self.assertFalse('<f>' in xml_of(response.content), 'a formula was written')
        found = [c for row in openpyxl.load_workbook(io.BytesIO(response.content)).active.iter_rows() for c in row
                 if isinstance(c.value, str) and 'evil.test' in c.value]
        self.assertTrue(found)
        self.assertTrue(all(c.data_type == 's' for c in found))

    def test_the_csv_report_defuses_a_formula_like_name(self):
        response = self.client.get('/api/reports/production/')
        self.assertEqual(response.status_code, 200)
        cells = [cell for row in csv.reader(io.StringIO(response.content.decode('utf-8-sig'))) for cell in row if 'evil.test' in cell]
        self.assertTrue(cells)
        self.assertTrue(all(not c.startswith('=') for c in cells), cells)
