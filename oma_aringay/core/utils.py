import csv
import datetime
import io
import math
from decimal import Decimal

from django.db.models import Q
from django.utils import timezone


def build_name_search_q(search, field_names):
    """Q object matching a person-name search across split first/middle/last
    name columns, correctly handling a full-name query like "Juan dela Cruz"
    that a plain single-field icontains against `search` as one string can
    never match (no single column holds the whole name together).

    Splits `search` on whitespace and requires EVERY resulting token to
    match at least one of `field_names` (icontains) — AND across tokens, OR
    across fields per token. A one-word search behaves exactly as a plain
    icontains across the given fields always did; "Juan dela Cruz" now
    additionally requires "Juan", "dela", and "Cruz" each match some name
    field, which a first_name="Juan"/last_name="Dela Cruz" row satisfies.

    Returns Q() (matches everything) for a blank search — callers should
    still guard the call with `if search:` for their own no-op-when-blank
    behavior, same as before this helper existed.
    """
    tokens = search.split()
    if not tokens:
        return Q()
    combined = Q()
    for token in tokens:
        token_q = Q()
        for field_name in field_names:
            token_q |= Q(**{f'{field_name}__icontains': token})
        combined &= token_q
    return combined


# ── Import-file safety ────────────────────────────────────────────────────
# Every Import File feature reads its upload through the helpers below, so a
# bad file (wrong encoding, corrupt/renamed workbook, a spreadsheet whose
# formatting runs to row 1,048,576, an enormous upload) gets ONE clear message
# instead of a server error or a hung request.
MAX_IMPORT_BYTES = 10 * 1024 * 1024   # the real 2026 RSBSA exports are ~1 MB each
MAX_IMPORT_ROWS = 20000               # ...and ~4,000 rows at most
_MAX_TRAILING_BLANK_ROWS = 1000       # a sheet "ends" after this many empty rows in a row


class ImportFileError(ValueError):
    """A user-facing 'this file can't be imported' problem. It is a
    ValueError so callers that already turn ValueError into a 400 with the
    message keep working unchanged."""


class UnsupportedFileType(ImportFileError):
    """Raised by iter_spreadsheet_rows for anything but .xlsx/.xls/.csv."""


def check_upload_size(fileobj):
    size = getattr(fileobj, 'size', None)
    if size is not None and size > MAX_IMPORT_BYTES:
        raise ImportFileError(
            f'The file is too large ({size / 1024 / 1024:.1f} MB). The limit is '
            f'{MAX_IMPORT_BYTES // (1024 * 1024)} MB — split it into smaller files.'
        )
    if size == 0:
        raise ImportFileError('The file is empty.')


def decode_csv_bytes(raw):
    """Decode an uploaded CSV. Excel's own "CSV" save is Windows-1252 (or
    UTF-16 via "Unicode Text"), not UTF-8 — names like "Nuñez"/"Peña" would
    otherwise crash a strict UTF-8 decode. UTF-8 (with or without BOM) is
    still tried first."""
    import codecs

    if isinstance(raw, str):
        return raw
    if raw.startswith(codecs.BOM_UTF16_LE) or raw.startswith(codecs.BOM_UTF16_BE):
        text = raw.decode('utf-16')
    else:
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = raw.decode('cp1252', errors='replace')
    if '\x00' in text:
        raise ImportFileError('This does not look like a text CSV file (it contains binary data).')
    return text


def csv_delimiter_for(text):
    """The delimiter Excel actually used — a semicolon or tab in some
    regional settings — picked from the header line."""
    first = text.split('\n', 1)[0]
    delimiter = max((',', ';', '\t'), key=first.count)
    return delimiter if first.count(delimiter) else ','


def csv_reader_for(text):
    return csv.reader(io.StringIO(text), delimiter=csv_delimiter_for(text))


def cell_text(value):
    """A spreadsheet cell as clean text: None -> '', 5.0 -> '5' (Excel stores
    whole numbers as floats), datetimes as ISO dates, everything else
    stripped."""
    if value is None:
        return ''
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    return str(value).strip()


def _row_is_blank(row):
    return not any(c is not None and str(c).strip() != '' for c in row)


def iter_spreadsheet_rows(fileobj, filename, *, min_row=1):
    """Yield row tuples (raw cell values) from an uploaded .xlsx, .xls, or
    .csv file, skipping every row before `min_row` (1-indexed) — the same
    convention openpyxl's own `ws.iter_rows(min_row=N, values_only=True)`
    already uses, so this is a drop-in replacement for it regardless of
    which of the three formats the upload actually is.

    Genuine Excel date cells — in EITHER .xlsx or .xls — come back as real
    `datetime` objects, matching openpyxl's own behavior, not a raw Excel
    serial number.

    Hardened for real-world uploads: size cap, a friendly ImportFileError for
    corrupt/renamed/password-protected workbooks, encoding detection for CSV,
    trailing empty (formatted-only) rows are never yielded, and a row cap.
    Blank rows in the MIDDLE are still yielded, so callers' row numbers
    always match the sheet.

    Raises UnsupportedFileType for any other extension.
    """
    name = (filename or '').lower()
    if not name.endswith(('.xlsx', '.xlsm', '.xls', '.csv')):
        raise UnsupportedFileType('Only .xlsx, .xls, or .csv files are supported.')
    check_upload_size(fileobj)

    def raw_rows():
        if name.endswith(('.xlsx', '.xlsm')):
            import zipfile

            import openpyxl
            try:
                fileobj.seek(0)
                wb = openpyxl.load_workbook(fileobj, read_only=True, data_only=True)
                ws = wb.active
                if not hasattr(ws, 'iter_rows'):
                    ws = next((s for s in wb.worksheets if hasattr(s, 'iter_rows')), None)
                if ws is None:
                    raise ImportFileError('The workbook has no worksheet with data.')
            except ImportFileError:
                raise
            except (zipfile.BadZipFile, KeyError, OSError, ValueError, openpyxl.utils.exceptions.InvalidFileException):
                raise ImportFileError(
                    'The file could not be read as an Excel (.xlsx) workbook. It may be corrupt, '
                    'password-protected, or not really an .xlsx file (e.g. a renamed CSV).'
                )
            yield from ws.iter_rows(min_row=min_row, values_only=True)
            return

        if name.endswith('.xls'):
            import xlrd
            try:
                fileobj.seek(0)
                book = xlrd.open_workbook(file_contents=fileobj.read())
                sheet = book.sheet_by_index(0)
            except Exception:
                raise ImportFileError(
                    'The file could not be read as a legacy Excel (.xls) workbook. It may be corrupt, '
                    'password-protected, or not really an .xls file.'
                )
            for row_idx in range(min_row - 1, sheet.nrows):
                values = []
                for col_idx in range(sheet.ncols):
                    cell = sheet.cell(row_idx, col_idx)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        values.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                    else:
                        values.append(cell.value)
                yield tuple(values)
            return

        fileobj.seek(0)
        text = decode_csv_bytes(fileobj.read())
        for i, row in enumerate(csv_reader_for(text), start=1):
            if i < min_row:
                continue
            yield tuple(row)

    pending_blank = []
    yielded = 0
    for row in raw_rows():
        if _row_is_blank(row):
            pending_blank.append(row)
            if len(pending_blank) > _MAX_TRAILING_BLANK_ROWS:
                return  # formatted-but-empty tail: the data has ended
            continue
        for blank in pending_blank:
            yielded += 1
            yield blank
        pending_blank = []
        yielded += 1
        if yielded > MAX_IMPORT_ROWS:
            raise ImportFileError(
                f'The file has more than {MAX_IMPORT_ROWS:,} rows. Split it into smaller files and import them one at a time.'
            )
        yield row


def read_template_rows(fileobj, filename, headers, *, sample_row=None, note_prefixes=(), max_rows=1000):
    """Data rows of one of this project's downloadable import templates, as
    a list of (sheet_row_number, row) — found by the header row's CONTENT,
    not by position. The old "data starts at row 3" rule silently dropped a
    first data row whenever staff overwrote the template's sample row 2 (the
    natural thing to do), and processed the template's own note line as data.

    - The header row must be in the first few rows and its first cell must be
      headers[0] (case/space-insensitive); this also stops a file from the
      wrong page's template (e.g. Fisherfolk assistance uploaded on the Crop
      Assistance page) from being half-read as something else.
    - The template's untouched sample row, its note line, and fully blank rows
      are skipped; every other row is returned with its real sheet row number.
    """
    def norm(v):
        return ' '.join(cell_text(v).lower().split())

    all_rows = list(iter_spreadsheet_rows(fileobj, filename))
    header_idx = next(
        (i for i, r in enumerate(all_rows[:6]) if r and norm(r[0]) == norm(headers[0])), None,
    )
    if header_idx is None:
        raise ImportFileError(
            f'Could not find the template header row (first column "{headers[0]}"). '
            f'Download the template from this page, fill it in, and upload that file.'
        )

    sample = tuple(norm(c) for c in sample_row) if sample_row else None
    data = []
    for idx in range(header_idx + 1, len(all_rows)):
        row = all_rows[idx]
        if _row_is_blank(row):
            continue
        first = cell_text(row[0]) if row else ''
        if any(first.startswith(p) for p in note_prefixes):
            continue
        if sample and tuple(norm(c) for c in row[:len(sample)]) == sample:
            continue
        data.append((idx + 1, row))
    # These imports save through the normal serializers (several queries per
    # row) inside ONE request; far beyond this the request would outlast any
    # proxy/browser timeout. Hand-filled templates never get near it.
    if len(data) > max_rows:
        raise ImportFileError(
            f'This import accepts up to {max_rows:,} rows at a time (the file has {len(data):,}). '
            f'Split it into smaller files and import them one after another.'
        )
    return data


def run_import_atomically(run):
    """Run an import that may create many rows as ONE all-or-nothing database
    transaction. These imports used to save row by row, so an unexpected
    failure part-way through left a half-imported file that had to be
    cleaned up by hand. Rows that are merely INVALID are still just reported
    (they never reach the database); this only guards the unexpected — e.g.
    two people importing the same organization at once, where the loser hit
    a raw server error after the winner's rows were already in."""
    from django.db import IntegrityError, transaction
    from rest_framework.response import Response

    try:
        with transaction.atomic():
            return run()
    except IntegrityError:
        return Response({
            'error': 'A record from this file was saved by someone else while it was importing '
                     '(a duplicate was created at the same moment). Nothing was imported — '
                     'upload the file again to see the updated results.',
        }, status=409)


def harvest_gate_status(start_date, expected_date):
    """
    Record Harvest (crop and fisheries alike) is gated until at least 50%
    of the record's expected growth/culture duration has elapsed since
    planting/stocking. Reuses the same start_date/expected_date pair that
    already drives expected_harvest_date and the Activity Suggestion
    Engine's progress display (CropType.days_to_harvest /
    SpeciesOption.culture_duration_days), rather than a second duration
    field or lookup.

    Returns (is_gated, message) — message is None when not gated. A record
    missing either date, or with a non-positive computed duration, is never
    gated (nothing to measure against).
    """
    if not start_date or not expected_date:
        return False, None
    total_days = (expected_date - start_date).days
    if total_days <= 0:
        return False, None
    halfway_days = total_days / 2
    elapsed_days = (timezone.localdate() - start_date).days
    if elapsed_days >= halfway_days:
        return False, None
    days_remaining = math.ceil(halfway_days - elapsed_days)
    return True, (
        f'Harvest can only be recorded once at least half the expected growth duration has '
        f'passed ({max(elapsed_days, 0)} of {math.ceil(halfway_days)} days so far). '
        f'Available in about {days_remaining} more day(s).'
    )


def day_range_bounds(date_str, end_of_day=False):
    """
    Parse a 'YYYY-MM-DD' query param into a timezone-aware datetime at the
    start (or end) of that local day.

    Avoids Django's `__date` lookup, which on MySQL compiles to
    DATE(CONVERT_TZ(...)) — CONVERT_TZ silently returns NULL (excluding the
    row) unless MySQL's named-timezone tables are loaded, which they aren't
    on a stock XAMPP install. Comparing plain datetimes with __gte/__lte
    sidesteps that dependency entirely.
    """
    try:
        day = datetime.datetime.strptime(date_str, '%Y-%m-%d').date()
    except (ValueError, TypeError):
        return None
    time_part = datetime.time.max if end_of_day else datetime.time.min
    return timezone.make_aware(datetime.datetime.combine(day, time_part))


# ── Yield Analysis: shared "Possible Factors" wording + aggregate math ─────
#
# Used by production/views.py, fisheries/views.py, and reports/views.py so
# the on-screen Yield Analysis tabs and the PDF/Excel exports never diverge —
# this project has been bitten before by the same logic quietly drifting
# apart across multiple copies (see the fisherfolk-search bug). Frontend has
# its own mirror of the wording rules in oma-frontend/src/utils/
# yieldAnalysis.js (JS can't import Python) — keep the two in sync by rule,
# not by literal string.

_YIELD_FACTOR_HINTS = {
    'crop': 'weather, pests, and soil',
    'fisheries': 'water quality, disease, and weather',
}
_YIELD_NO_SCHEDULE_NOUN = {
    'crop': 'crop',
    'fisheries': 'species',
}


def possible_factors_text(variance_pct, missed_high_priority, activities_completed_count,
                           activities_total_count, domain='crop'):
    """
    "Possible Factors" text for one Yield Analysis record. Causal-adjacent
    "Missed: ..." wording is reserved for the one case it's actually
    informative for — a real shortfall (variance_pct < 0) alongside logged
    activity gaps. Showing that same wording next to a positive or unknown
    variance would imply the gaps caused a problem that didn't happen (or
    can't even be measured), so those cases get neutral, non-causal text
    instead.
    """
    if activities_total_count == 0:
        noun = _YIELD_NO_SCHEDULE_NOUN.get(domain, 'record')
        return f'No activity schedule defined for this {noun}.'

    has_gaps = bool(missed_high_priority)

    if variance_pct is None:
        if not has_gaps:
            return 'All key activities completed'
        return (
            f'Recommended activities logged: {activities_completed_count}/{activities_total_count} '
            '— no yield estimate available for comparison.'
        )

    if variance_pct >= 0:
        if not has_gaps:
            return 'All key activities completed'
        return (
            f'Recommended activities logged: {activities_completed_count}/{activities_total_count} '
            '— yield met or exceeded the estimate despite limited logged activity.'
        )

    # variance_pct < 0
    if not has_gaps:
        hint = _YIELD_FACTOR_HINTS.get(domain, 'external, unlogged conditions')
        return (
            'Most recommended activities were completed; the shortfall may reflect '
            f"factors outside this system's scope ({hint})."
        )
    parts = [f"{m['activity']} (Day {m['scheduled_day']})" for m in missed_high_priority]
    return 'Missed: ' + ', '.join(parts)


def new_yield_group():
    """A fresh accumulator for one Yield Analysis aggregate group (barangay/
    crop type/species). Pair with accumulate_yield_group()."""
    return {
        'count': 0,
        'total_loss_count': 0,
        'estimated_sum': Decimal('0'), 'estimated_n': 0,
        'actual_sum': Decimal('0'),
        'variance_pct_sum': 0.0, 'variance_pct_n': 0,
        'completion_pct_sum': 0.0, 'completion_pct_n': 0,
    }


def accumulate_yield_group(g, a):
    """
    Fold one analyzed Yield Analysis record (the dict returned by
    _analyze_crop/_analyze_stocking) into a group accumulator from
    new_yield_group().

    Total Loss records are counted (g['count'], g['total_loss_count']) and
    included in avg_actual_kg (their actual output really is 0, which is
    real information), but excluded from avg_variance_pct — a -100% Total
    Loss variance blended unweighted into a barangay average silently
    dominates it and gives a misleading read on how the "normal" harvests in
    that group performed. total_loss_count on the result makes the
    exclusion visible instead of silent.
    """
    g['count'] += 1
    if a['is_total_loss']:
        g['total_loss_count'] += 1
    if a['estimated_kg'] is not None:
        g['estimated_sum'] += Decimal(str(a['estimated_kg']))
        g['estimated_n'] += 1
    g['actual_sum'] += Decimal(str(a['actual_total_kg']))
    if a['variance_pct'] is not None and not a['is_total_loss']:
        g['variance_pct_sum'] += a['variance_pct']
        g['variance_pct_n'] += 1
    if a['activities_total_count'] > 0:
        g['completion_pct_sum'] += (a['activities_completed_count'] / a['activities_total_count']) * 100
        g['completion_pct_n'] += 1
