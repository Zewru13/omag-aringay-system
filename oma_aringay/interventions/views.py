import uuid

from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import transaction
from django.shortcuts import get_object_or_404
from django.utils import timezone
from rest_framework import generics
from rest_framework import status as drf_status
from rest_framework.exceptions import ValidationError
from rest_framework.parsers import FormParser, JSONParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from accounts.permissions import (
    IsAdminOrEvaluator, IsCitizen, IsOrgMemberRecordCitizen, IsStaffOrFisherfolk,
    is_fisherfolk_citizen, scope_to_fisherfolk,
)
from accounts.utils import accountless_or_unactivated_ids
from associations.models import FarmersOrganizationMember, FisherfolkOrganizationMember
from core.notifications import notify_staff, notify_user, resolve_owner_user
from core.photos import discard_uploaded_files, validated_photos
from core.validators import validate_uploaded_image

from .lifecycle import advance_due_distributions, advance_due_distributions_for_org, notify_assistance_recorded
from .models import (
    BeneficiaryConfirmationRequestPhoto, BeneficiaryRecordPhoto, FisherfolkBeneficiaryConfirmationRequestPhoto,
    BeneficiaryConfirmationRequest,
    BeneficiaryRecord,
    DistributionStatus,
    FisherfolkBeneficiaryConfirmationRequest,
    FisherfolkBeneficiaryRecord,
)
from .distributions import build_distributions, filter_records, remove_records
from .serializers import (
    AccountlessConfirmableSerializer,
    BeneficiaryConfirmationRequestSerializer,
    BeneficiaryRecordCreateSerializer,
    BeneficiaryRecordUpdateSerializer,
    FisherfolkAccountlessConfirmableSerializer,
    FisherfolkBeneficiaryConfirmationRequestSerializer,
    FisherfolkBeneficiaryRecordCreateSerializer,
    FisherfolkBeneficiaryRecordUpdateSerializer,
)


class PendingOnlyEditMixin:
    """
    Who may still edit an assistance record, by its state:
      - an organization (CITIZEN account) only while it is PENDING: once the record is due it belongs to the recipient;
      - staff while it is PENDING, AWAITING_CONFIRMATION (new assistance starts there, so a typo can still be fixed) or
        CONFIRMED, but not once the recipient reported non-receipt or the case was resolved.
    """
    ORGANIZATION_EDITABLE = (DistributionStatus.PENDING,)
    STAFF_EDITABLE = (
        DistributionStatus.PENDING, DistributionStatus.AWAITING_CONFIRMATION, DistributionStatus.CONFIRMED,
    )

    def perform_update(self, serializer):
        is_organization = self.request.user.role == 'CITIZEN'
        allowed = self.ORGANIZATION_EDITABLE if is_organization else self.STAFF_EDITABLE
        if serializer.instance.status not in allowed:
            raise ValidationError(
                f'This record cannot be edited — it is {serializer.instance.get_status_display()}, '
                + ('not Pending.' if is_organization else 'not Pending, Awaiting Confirmation or Confirmed.')
            )
        serializer.save()


class BeneficiaryRecordDetailView(PendingOnlyEditMixin, generics.RetrieveUpdateAPIView):
    """
    Batch 12: GET/PATCH by id for any BeneficiaryRecord, shared by staff
    (unrestricted, via IsOrgMemberRecordCitizen's ADMIN/EVALUATOR bypass)
    and ORGANIZATION citizens (own-submitted, own-member entries only).
    """
    queryset = BeneficiaryRecord.objects.select_related('farmer')
    serializer_class = BeneficiaryRecordUpdateSerializer
    permission_classes = [IsAuthenticated, IsOrgMemberRecordCitizen]
    org_linked_record_type = 'FARMERS_ORG'   # a farmer id here must never be matched by a fisherfolk organization
    http_method_names = ['get', 'patch']
    lookup_field = 'id'

    def get_object_person_id(self, obj):
        return obj.farmer_id

    def get_object_creator_id(self, obj):
        return obj.recorded_by_id


class BeneficiaryAdHocListCreateView(generics.ListCreateAPIView):
    serializer_class = BeneficiaryRecordCreateSerializer
    pagination_class = None
    org_linked_record_type = 'FARMERS_ORG'

    def get_permissions(self):
        # Batch 12: LIST stays staff-only — an ORGANIZATION citizen browsing
        # every record here would see other organizations'/staff's
        # submissions for shared members. CREATE opens to them via
        # IsOrgMemberRecordCitizen, which staff pass through unconditionally
        # either way.
        if self.request.method == 'GET':
            return [IsAuthenticated(), IsAdminOrEvaluator()]
        return [IsAuthenticated(), IsOrgMemberRecordCitizen()]

    def get_target_person_id(self, request):
        return request.data.get('farmer')

    def get_queryset(self):
        qs = BeneficiaryRecord.objects.select_related('farmer', 'recorded_by').prefetch_related('photos').order_by('-date_given', '-id')
        params = self.request.query_params
        if params.get('farmer'):
            qs = qs.filter(farmer_id=params['farmer'])
        if params.get('year'):
            qs = qs.filter(year=params['year'])
        if params.get('barangay'):
            qs = qs.filter(farmer__barangay=params['barangay'])
        # Step 5: "Needs Follow-up" — ?status=NOT_RECEIVED on the same
        # existing list endpoint CropAssistancePage already calls, rather
        # than a separate view.
        if params.get('status'):
            qs = qs.filter(status=params['status'])
        return qs

    def perform_create(self, serializer):
        record = serializer.save(recorded_by=self.request.user, distribution_batch=uuid.uuid4())
        notify_assistance_recorded(record, self.request.user)


class BeneficiaryAdHocDetailView(PendingOnlyEditMixin, generics.RetrieveUpdateDestroyAPIView):
    serializer_class = BeneficiaryRecordCreateSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    queryset = BeneficiaryRecord.objects.select_related('farmer')
    lookup_field = 'id'
    http_method_names = ['get', 'patch', 'delete']

    def destroy(self, request, *args, **kwargs):
        if request.user.role != 'ADMIN':
            return Response(
                {'detail': 'Only administrators can delete assistance records.'},
                status=drf_status.HTTP_403_FORBIDDEN,
            )
        return super().destroy(request, *args, **kwargs)


class FisherfolkBeneficiaryRecordDetailView(PendingOnlyEditMixin, generics.RetrieveUpdateAPIView):
    """Fisherfolk equivalent of BeneficiaryRecordDetailView."""
    queryset = FisherfolkBeneficiaryRecord.objects.select_related('fisherfolk')
    serializer_class = FisherfolkBeneficiaryRecordUpdateSerializer
    permission_classes = [IsAuthenticated, IsOrgMemberRecordCitizen]
    org_linked_record_type = 'FISHERFOLK_ORG'   # a fisherfolk id here must never be matched by a farmers' organization
    http_method_names = ['get', 'patch']
    lookup_field = 'id'

    def get_object_person_id(self, obj):
        return obj.fisherfolk_id

    def get_object_creator_id(self, obj):
        return obj.recorded_by_id


class FisherfolkBeneficiaryAdHocListCreateView(generics.ListCreateAPIView):
    serializer_class = FisherfolkBeneficiaryRecordCreateSerializer
    pagination_class = None
    org_linked_record_type = 'FISHERFOLK_ORG'

    def get_permissions(self):
        if self.request.method == 'GET':
            # The Admin Fisheries Assistance page is reused by fisherfolk accounts read-only: a fisherfolk
            # citizen lists only their OWN assistance (scoped below). Creating stays staff/organization.
            return [IsAuthenticated(), IsStaffOrFisherfolk()]
        return [IsAuthenticated(), IsOrgMemberRecordCitizen()]

    def get_target_person_id(self, request):
        return request.data.get('fisherfolk')

    def get_queryset(self):
        if is_fisherfolk_citizen(self.request.user):
            advance_due_distributions(self.request.user)   # same lazy status step as MyFisherfolkBeneficiaryHistoryView
        qs = scope_to_fisherfolk(FisherfolkBeneficiaryRecord.objects.all(), self.request.user) \
            .select_related('fisherfolk', 'recorded_by').order_by('-date_given', '-id')
        params = self.request.query_params
        if params.get('fisherfolk'):
            qs = qs.filter(fisherfolk_id=params['fisherfolk'])
        if params.get('year'):
            qs = qs.filter(year=params['year'])
        if params.get('barangay'):
            qs = qs.filter(fisherfolk__barangay=params['barangay'])
        if params.get('status'):
            qs = qs.filter(status=params['status'])
        return qs

    def perform_create(self, serializer):
        record = serializer.save(recorded_by=self.request.user, distribution_batch=uuid.uuid4())
        notify_assistance_recorded(record, self.request.user)


class FisherfolkBeneficiaryAdHocDetailView(PendingOnlyEditMixin, generics.RetrieveUpdateDestroyAPIView):
    serializer_class = FisherfolkBeneficiaryRecordCreateSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    queryset = FisherfolkBeneficiaryRecord.objects.select_related('fisherfolk')
    lookup_field = 'id'
    http_method_names = ['get', 'patch', 'delete']

    def destroy(self, request, *args, **kwargs):
        if request.user.role != 'ADMIN':
            return Response(
                {'detail': 'Only administrators can delete assistance records.'},
                status=drf_status.HTTP_403_FORBIDDEN,
            )
        return super().destroy(request, *args, **kwargs)


class MyBeneficiaryHistoryView(generics.ListAPIView):
    """
    Batch 11: CITIZEN-only, read-only view of the caller's own assistance
    history, resolved via get_linked_record() (never a farmer id from the
    request), so a citizen cannot pass someone else's id to see their
    history.
    """
    serializer_class = BeneficiaryRecordCreateSerializer
    permission_classes = [IsAuthenticated, IsCitizen]
    pagination_class = None

    def get_queryset(self):
        # Batch 12: explicit type guard — an ORGANIZATION-linked citizen's
        # linked_record_id points into a different table entirely, and
        # comparing it against farmer_id here would risk an accidental
        # int collision. This view is FARMER-only; see
        # MyOrgBeneficiaryHistoryView for the ORGANIZATION-role equivalent
        # (different semantics: recorded_by=self, not farmer=self — an
        # organization is never a beneficiary subject itself).
        advance_due_distributions(self.request.user)
        farmer_id = self.request.user.get_linked_farmer_id()
        if farmer_id is None:
            return BeneficiaryRecord.objects.none()
        return BeneficiaryRecord.objects.filter(
            farmer_id=farmer_id,
        ).select_related('farmer', 'recorded_by').prefetch_related('photos').order_by('-date_given', '-id')


class MyFisherfolkBeneficiaryHistoryView(generics.ListAPIView):
    """Batch 11: fisherfolk equivalent of MyBeneficiaryHistoryView."""
    serializer_class = FisherfolkBeneficiaryRecordCreateSerializer
    permission_classes = [IsAuthenticated, IsCitizen]
    pagination_class = None

    def get_queryset(self):
        # Batch 12: see MyBeneficiaryHistoryView's comment — same type-guard
        # fix, FISHERFOLK-only here.
        advance_due_distributions(self.request.user)
        fisherfolk_id = self.request.user.get_linked_fisherfolk_id()
        if fisherfolk_id is None:
            return FisherfolkBeneficiaryRecord.objects.none()
        return FisherfolkBeneficiaryRecord.objects.filter(
            fisherfolk_id=fisherfolk_id,
        ).select_related('fisherfolk', 'recorded_by').order_by('-date_given', '-id')


class MyOrgBeneficiaryHistoryView(generics.ListAPIView):
    """
    Batch 12: an ORGANIZATION citizen's own reporting history — records
    THEY submitted (recorded_by = themselves). Deliberately NOT "every
    record touching this org's members" — an organization is never a
    beneficiary subject itself, and surfacing other organizations'/staff's
    submissions for the same member is explicitly out of this batch's
    scope (Step 4). Different filter shape from MyBeneficiaryHistoryView
    on purpose: that one is farmer_id=self (person is the subject), this
    one is recorded_by=self (org is the reporter).
    """
    serializer_class = BeneficiaryRecordCreateSerializer
    permission_classes = [IsAuthenticated, IsCitizen]
    pagination_class = None

    def get_queryset(self):
        if self.request.user.linked_record_type != 'FARMERS_ORG':
            return BeneficiaryRecord.objects.none()
        return BeneficiaryRecord.objects.filter(
            recorded_by=self.request.user,
        ).select_related('farmer').prefetch_related('photos').order_by('-date_given', '-id')


class MyOrgFisherfolkBeneficiaryHistoryView(generics.ListAPIView):
    """Batch 12: fisherfolk-organization equivalent of MyOrgBeneficiaryHistoryView."""
    serializer_class = FisherfolkBeneficiaryRecordCreateSerializer
    permission_classes = [IsAuthenticated, IsCitizen]
    pagination_class = None

    def get_queryset(self):
        if self.request.user.linked_record_type != 'FISHERFOLK_ORG':
            return FisherfolkBeneficiaryRecord.objects.none()
        return FisherfolkBeneficiaryRecord.objects.filter(
            recorded_by=self.request.user,
        ).select_related('fisherfolk').order_by('-date_given', '-id')


class MyOrgDistributionsView(APIView):
    """An organization's own assistance history as distributions (one entry per distribution, expandable to its
    members), with the filters `status` (a status or `needs-follow-up`), `assistance_type`, `date_from`, `date_to`,
    `member` and `source_agency`. Only records this organization account recorded, never another organization's."""
    permission_classes = [IsAuthenticated, IsCitizen]
    org_linked_record_type = 'FARMERS_ORG'
    record_model = BeneficiaryRecord
    person_field = 'farmer'

    def get(self, request):
        if request.user.linked_record_type != self.org_linked_record_type:
            return Response([])
        queryset = self.record_model.objects.filter(recorded_by=request.user).select_related(self.person_field)
        records = filter_records(queryset, self.person_field, request.query_params)
        return Response(build_distributions(records, self.person_field))


class MyOrgFisherfolkDistributionsView(MyOrgDistributionsView):
    org_linked_record_type = 'FISHERFOLK_ORG'
    record_model = FisherfolkBeneficiaryRecord
    person_field = 'fisherfolk'


class AssistanceRemovalView(APIView):
    """Take entries out of the list before the recipient confirms them: body {domain: 'crop'|'fisheries', record_ids: [...]}
    or {domain, distribution: <group id>}. Staff may remove any entry; an organization account only what it recorded."""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        user = request.user
        domain = request.data.get('domain')
        denied = Response({'detail': 'You do not have permission to perform this action.'}, status=403)
        if user.role == 'CITIZEN':
            expected = {'crop': 'FARMERS_ORG', 'fisheries': 'FISHERFOLK_ORG'}.get(domain)
            if expected is None or user.linked_record_type != expected:
                return denied
        elif user.role not in ('ADMIN', 'EVALUATOR'):
            return denied
        record_ids = request.data.get('record_ids')
        if record_ids is not None and not isinstance(record_ids, list):
            raise ValidationError({'record_ids': ['Send a list of record ids.']})
        try:
            result = remove_records(domain, user, record_ids=record_ids, distribution=request.data.get('distribution'))
        except (TypeError, ValueError):
            raise ValidationError({'record_ids': ['Send a list of record ids.']})
        return Response(result)


# ─── Scheduled distribution + confirmation lifecycle: recipient & staff actions ──
# Structural reference: production.views.CropDisasterResolveView — plain
# APIView (no queryset for DRF's object-level permission machinery to hook
# into), explicit get_object_or_404, explicit status guard, manual
# ownership check for the CITIZEN branch rather than a permission class,
# same as that view's own documented reasoning.

def _beneficiary_label(record, person, assistance_type=None):
    return f'{assistance_type or record.assistance_type or "Assistance"} — {person.full_name}'


def _check_recipient_permission(request, owner_linked_type, owner_id):
    """Shared by both domains' Confirm/ReportNotReceived views. Own-record-
    only for the CITIZEN who IS the recipient — not the reporting
    organization (IsOrgMemberRecordCitizen is a different, unrelated
    permission for a different actor). Staff may act too, same as disaster
    resolution's own precedent."""
    if request.user.role in ('ADMIN', 'EVALUATOR'):
        return None
    if request.user.role != 'CITIZEN':
        return Response({'detail': 'You do not have permission to perform this action.'}, status=403)
    if request.user.linked_record_type != owner_linked_type or request.user.linked_record_id != owner_id:
        return Response({'detail': 'You do not have permission to perform this action.'}, status=403)
    return None


class BeneficiaryConfirmReceivedView(APIView):
    permission_classes = [IsAuthenticated]
    # Multipart so the recipient can attach an optional proof photo; JSON still
    # works for the bare confirm the older clients send.
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def post(self, request, pk):
        record = get_object_or_404(BeneficiaryRecord.objects.select_related('farmer'), pk=pk)
        denied = _check_recipient_permission(request, 'FARMER', record.farmer_id)
        if denied:
            return denied
        if record.status != DistributionStatus.AWAITING_CONFIRMATION:
            return Response(
                {'detail': f'This record is not awaiting confirmation (current status: {record.get_status_display()}).'},
                status=400,
            )
        note = (request.data.get('confirmation_note') or '').strip()
        # 0-10 proof photos (the older single `confirmation_photo` counts as one), all validated before anything
        # changes; the status change and the photo rows are one transaction and a failed attempt's files are removed.
        photos = validated_photos(
            request, 'confirmation_photos', legacy_field='confirmation_photo', limit=BeneficiaryRecordPhoto.MAX_PER_PARENT)
        written = []
        try:
            with transaction.atomic():
                # Re-read under a row lock: a second simultaneous confirmation (double tap, retry, two tabs) waits
                # here and then finds the record already confirmed instead of adding its photos again.
                record = BeneficiaryRecord.objects.select_for_update().get(pk=record.pk)
                if record.status != DistributionStatus.AWAITING_CONFIRMATION:
                    return Response(
                        {'detail': f'This record is not awaiting confirmation (current status: {record.get_status_display()}).'},
                        status=400,
                    )
                record.confirmation_note = note
                record.status = DistributionStatus.CONFIRMED
                record.confirmed_at = timezone.now()
                record.save()
                for f in photos:
                    row = BeneficiaryRecordPhoto.objects.create(beneficiary_record=record, image=f, uploaded_by=request.user)
                    written.append(row.image.name)
                record = BeneficiaryRecord.objects.select_related('farmer').prefetch_related('photos').get(pk=record.pk)
                data = BeneficiaryRecordCreateSerializer(record, context={'request': request}).data
        except Exception:
            discard_uploaded_files(written)
            raise
        return Response(data)


class BeneficiaryReportNotReceivedView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        record = get_object_or_404(BeneficiaryRecord.objects.select_related('farmer', 'attributed_organization'), pk=pk)
        denied = _check_recipient_permission(request, 'FARMER', record.farmer_id)
        if denied:
            return denied
        if record.status != DistributionStatus.AWAITING_CONFIRMATION:
            return Response(
                {'detail': f'This record is not awaiting confirmation (current status: {record.get_status_display()}).'},
                status=400,
            )
        reason = (request.data.get('non_receipt_reason') or '').strip()
        if not reason:
            return Response({'detail': 'A reason is required.'}, status=400)

        record.status = DistributionStatus.NOT_RECEIVED
        record.non_receipt_reason = reason
        record.save()

        label = _beneficiary_label(record, record.farmer)
        notify_staff(
            actor=request.user, verb='reported not receiving', record_type='beneficiary',
            record_id=record.id, record_label=label, detail=reason,
        )
        notify_user(
            recipient=resolve_owner_user(farmers_organization_id=record.attributed_organization_id),
            actor=request.user, verb='reported not receiving', record_type='beneficiary',
            record_id=record.id, record_label=label, detail=reason,
        )
        return Response(BeneficiaryRecordCreateSerializer(record).data)


class BeneficiaryMarkResolvedView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def post(self, request, pk):
        record = get_object_or_404(BeneficiaryRecord, pk=pk)
        if record.status != DistributionStatus.NOT_RECEIVED:
            return Response(
                {'detail': f'This record is not a Not Received case (current status: {record.get_status_display()}).'},
                status=400,
            )
        notes = (request.data.get('follow_up_notes') or '').strip()
        if not notes:
            return Response({'detail': 'Follow-up notes are required.'}, status=400)
        record.status = DistributionStatus.RESOLVED
        record.follow_up_notes = notes
        record.save()
        return Response(BeneficiaryRecordCreateSerializer(record).data)


class FisherfolkBeneficiaryConfirmReceivedView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        record = get_object_or_404(FisherfolkBeneficiaryRecord.objects.select_related('fisherfolk'), pk=pk)
        denied = _check_recipient_permission(request, 'FISHERFOLK', record.fisherfolk_id)
        if denied:
            return denied
        if record.status != DistributionStatus.AWAITING_CONFIRMATION:
            return Response(
                {'detail': f'This record is not awaiting confirmation (current status: {record.get_status_display()}).'},
                status=400,
            )
        record.status = DistributionStatus.CONFIRMED
        record.confirmed_at = timezone.now()
        record.save()
        return Response(FisherfolkBeneficiaryRecordCreateSerializer(record).data)


class FisherfolkBeneficiaryReportNotReceivedView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        record = get_object_or_404(
            FisherfolkBeneficiaryRecord.objects.select_related('fisherfolk', 'attributed_fisherfolk_organization'), pk=pk,
        )
        denied = _check_recipient_permission(request, 'FISHERFOLK', record.fisherfolk_id)
        if denied:
            return denied
        if record.status != DistributionStatus.AWAITING_CONFIRMATION:
            return Response(
                {'detail': f'This record is not awaiting confirmation (current status: {record.get_status_display()}).'},
                status=400,
            )
        reason = (request.data.get('non_receipt_reason') or '').strip()
        if not reason:
            return Response({'detail': 'A reason is required.'}, status=400)

        record.status = DistributionStatus.NOT_RECEIVED
        record.non_receipt_reason = reason
        record.save()

        label = _beneficiary_label(record, record.fisherfolk)
        notify_staff(
            actor=request.user, verb='reported not receiving', record_type='ff_beneficiary',
            record_id=record.id, record_label=label, detail=reason,
        )
        notify_user(
            recipient=resolve_owner_user(fisherfolk_organization_id=record.attributed_fisherfolk_organization_id),
            actor=request.user, verb='reported not receiving', record_type='ff_beneficiary',
            record_id=record.id, record_label=label, detail=reason,
        )
        return Response(FisherfolkBeneficiaryRecordCreateSerializer(record).data)


class FisherfolkBeneficiaryMarkResolvedView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def post(self, request, pk):
        record = get_object_or_404(FisherfolkBeneficiaryRecord, pk=pk)
        if record.status != DistributionStatus.NOT_RECEIVED:
            return Response(
                {'detail': f'This record is not a Not Received case (current status: {record.get_status_display()}).'},
                status=400,
            )
        notes = (request.data.get('follow_up_notes') or '').strip()
        if not notes:
            return Response({'detail': 'Follow-up notes are required.'}, status=400)
        record.status = DistributionStatus.RESOLVED
        record.follow_up_notes = notes
        record.save()
        return Response(FisherfolkBeneficiaryRecordCreateSerializer(record).data)


# ─── Accountless-beneficiary confirmation ────────────────────────────────────
# See interventions.models.BeneficiaryConfirmationRequest's own module
# comment for the full design reasoning (separate model, not new
# BeneficiaryRecord fields; mirrored per-domain, not one polymorphic model;
# additive to DistributionStatus, not a new state in it).

def _check_org_confirms_on_behalf_permission(request, org_linked_type, person_id):
    """
    Only a CITIZEN linked to `org_linked_type`, whose organization is the
    ACTUAL current organization of `person_id`, may submit an on-behalf-of
    confirmation. Deliberately NO ADMIN/EVALUATOR bypass, unlike
    _check_recipient_permission's staff bypass above — staff already has an
    unmediated way to confirm receipt directly (BeneficiaryConfirmReceivedView
    lets ADMIN/EVALUATOR act as the recipient with no approval loop), so this
    feature exists specifically to let an ORGANIZATION do what only staff
    could do before, subject to admin review because the org is not the
    actual recipient.
    """
    if request.user.role != 'CITIZEN' or request.user.linked_record_type != org_linked_type:
        return Response({'detail': 'You do not have permission to perform this action.'}, status=403)
    org_id = request.user.linked_record_id
    if org_linked_type == 'FARMERS_ORG':
        is_member = FarmersOrganizationMember.objects.filter(
            farmer_id=person_id, organization_id=org_id, is_active=True,
        ).exists()
    else:
        is_member = FisherfolkOrganizationMember.objects.filter(
            fisherfolk_id=person_id, organization_id=org_id, is_active=True,
        ).exists()
    if not is_member:
        return Response({'detail': 'You do not have permission to perform this action.'}, status=403)
    return None


# An organization vouches for a member it confirms for, so it must show proof: at least this many photos.
MIN_ON_BEHALF_PHOTOS = 2
# The most characters a rejection note can hold (BeneficiaryConfirmationRequest.rejection_reason).
REJECTION_NOTE_MAX = 255


def _require_proof_photos(photos):
    if len(photos) < MIN_ON_BEHALF_PHOTOS:
        raise ValidationError({'photos': [f'Add at least {MIN_ON_BEHALF_PHOTOS} photos as proof.']})


def _required_rejection_note(request):
    """The note staff must write when they reject an on-behalf-of confirmation: it is sent to the organization,
    so a blank one, or one too long to store whole, is refused rather than saved empty or cut short."""
    note = str(request.data.get('rejection_reason') or '').strip()
    if not note:
        raise ValidationError({'rejection_reason': ['Write why you are rejecting this confirmation, so the organization knows what to fix.']})
    if len(note) > REJECTION_NOTE_MAX:
        raise ValidationError({'rejection_reason': [f'Keep the note to {REJECTION_NOTE_MAX} characters or fewer.']})
    return note


class MyOrgAccountlessConfirmablesView(generics.ListAPIView):
    """
    An ORGANIZATION citizen's confirmable-on-behalf-of queue: BeneficiaryRecord
    rows that are AWAITING_CONFIRMATION for a CURRENT member of this org whose
    citizen account genuinely has no way to self-confirm (no shell, or one
    that has never been activated — accounts.utils.accountless_or_unactivated_ids,
    the same signal driving the "Activate Account" button elsewhere).

    Deliberately NOT a widening of MyOrgBeneficiaryHistoryView's scope above
    (that view's own docstring explicitly scopes surfacing other reporters'
    submissions as "Step 4, out of scope") — this is a new, narrowly-purposed
    endpoint instead, needed because most real distributions run through
    staff's bulk Distribute-to-Organization flow, where recorded_by is the
    staff user, not the org — MyOrgBeneficiaryHistoryView's recorded_by=self
    filter would never surface those rows to the org at all.
    """
    serializer_class = AccountlessConfirmableSerializer
    permission_classes = [IsAuthenticated, IsCitizen]
    pagination_class = None

    def get_queryset(self):
        if self.request.user.linked_record_type != 'FARMERS_ORG':
            return BeneficiaryRecord.objects.none()
        org_id = self.request.user.linked_record_id
        # Lazily advances any of THIS org's accountless members' due
        # PENDING records to AWAITING_CONFIRMATION first — see
        # advance_due_distributions_for_org's own docstring for why this
        # can never happen via the normal citizen-side path for them.
        advance_due_distributions_for_org('FARMERS_ORG', org_id)
        candidates = BeneficiaryRecord.objects.filter(
            status=DistributionStatus.AWAITING_CONFIRMATION,
            farmer__farmer_organizations__organization_id=org_id,
            farmer__farmer_organizations__is_active=True,
        ).select_related('farmer').prefetch_related('confirmation_requests').distinct()
        farmer_ids = [r.farmer_id for r in candidates]
        eligible_ids = accountless_or_unactivated_ids('FARMER', farmer_ids)
        return candidates.filter(farmer_id__in=eligible_ids).order_by('-date_given', '-id')


class MyOrgFisherfolkAccountlessConfirmablesView(generics.ListAPIView):
    """Fisherfolk equivalent of MyOrgAccountlessConfirmablesView above."""
    serializer_class = FisherfolkAccountlessConfirmableSerializer
    permission_classes = [IsAuthenticated, IsCitizen]
    pagination_class = None

    def get_queryset(self):
        if self.request.user.linked_record_type != 'FISHERFOLK_ORG':
            return FisherfolkBeneficiaryRecord.objects.none()
        org_id = self.request.user.linked_record_id
        advance_due_distributions_for_org('FISHERFOLK_ORG', org_id)
        candidates = FisherfolkBeneficiaryRecord.objects.filter(
            status=DistributionStatus.AWAITING_CONFIRMATION,
            fisherfolk__fisherfolk_organizations__organization_id=org_id,
            fisherfolk__fisherfolk_organizations__is_active=True,
        ).select_related('fisherfolk').prefetch_related('confirmation_requests').distinct()
        fisherfolk_ids = [r.fisherfolk_id for r in candidates]
        eligible_ids = accountless_or_unactivated_ids('FISHERFOLK', fisherfolk_ids)
        return candidates.filter(fisherfolk_id__in=eligible_ids).order_by('-date_given', '-id')


class BeneficiaryConfirmOnBehalfView(APIView):
    """
    An ORGANIZATION confirms receipt on behalf of a member who genuinely has
    no working citizen account. Creates a PENDING BeneficiaryConfirmationRequest
    for admin review rather than writing CONFIRMED directly — the org is not
    the actual recipient, so this goes through the same pending-then-staff-
    approval shape core.models.ProfileChangeRequest already established, not
    straight to BeneficiaryConfirmReceivedView's immediate-confirm behavior.
    """
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser]

    def post(self, request, pk):
        record = get_object_or_404(BeneficiaryRecord.objects.select_related('farmer'), pk=pk)
        denied = _check_org_confirms_on_behalf_permission(request, 'FARMERS_ORG', record.farmer_id)
        if denied:
            return denied
        if record.status != DistributionStatus.AWAITING_CONFIRMATION:
            return Response(
                {'detail': f'This record is not awaiting confirmation (current status: {record.get_status_display()}).'},
                status=400,
            )
        eligible_ids = accountless_or_unactivated_ids('FARMER', [record.farmer_id])
        if record.farmer_id not in eligible_ids:
            return Response(
                {'detail': 'This beneficiary has a working account and must confirm receipt themselves.'},
                status=400,
            )

        photos = validated_photos(request, 'photos', legacy_field='photo', limit=BeneficiaryConfirmationRequestPhoto.MAX_PER_PARENT)
        _require_proof_photos(photos)

        written = []
        try:
            with transaction.atomic():
                existing = list(
                    BeneficiaryConfirmationRequest.objects.select_for_update().filter(
                        beneficiary_record=record, status=BeneficiaryConfirmationRequest.Status.PENDING,
                    )
                )
                if existing:
                    return Response(
                        {'detail': 'A confirmation request for this record is already pending admin review.'},
                        status=400,
                    )
                req = BeneficiaryConfirmationRequest.objects.create(
                    beneficiary_record=record,
                    submitted_by_organization_id=request.user.linked_record_id,
                    submitted_by=request.user,
                )
                for f in photos:
                    row = BeneficiaryConfirmationRequestPhoto.objects.create(confirmation_request=req, image=f, uploaded_by=request.user)
                    written.append(row.image.name)
        except Exception:
            discard_uploaded_files(written)
            raise

        notify_staff(
            actor=request.user, verb='submitted an on-behalf-of confirmation for',
            record_type='beneficiary_obo', record_id=req.id,
            record_label=_beneficiary_label(record, record.farmer),
        )
        return Response(BeneficiaryConfirmationRequestSerializer(req, context={'request': request}).data, status=drf_status.HTTP_201_CREATED)


class FisherfolkBeneficiaryConfirmOnBehalfView(APIView):
    """Fisherfolk equivalent of BeneficiaryConfirmOnBehalfView above."""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser]

    def post(self, request, pk):
        record = get_object_or_404(FisherfolkBeneficiaryRecord.objects.select_related('fisherfolk'), pk=pk)
        denied = _check_org_confirms_on_behalf_permission(request, 'FISHERFOLK_ORG', record.fisherfolk_id)
        if denied:
            return denied
        if record.status != DistributionStatus.AWAITING_CONFIRMATION:
            return Response(
                {'detail': f'This record is not awaiting confirmation (current status: {record.get_status_display()}).'},
                status=400,
            )
        eligible_ids = accountless_or_unactivated_ids('FISHERFOLK', [record.fisherfolk_id])
        if record.fisherfolk_id not in eligible_ids:
            return Response(
                {'detail': 'This beneficiary has a working account and must confirm receipt themselves.'},
                status=400,
            )

        photos = validated_photos(request, 'photos', legacy_field='photo', limit=FisherfolkBeneficiaryConfirmationRequestPhoto.MAX_PER_PARENT)
        _require_proof_photos(photos)

        written = []
        try:
            with transaction.atomic():
                existing = list(
                    FisherfolkBeneficiaryConfirmationRequest.objects.select_for_update().filter(
                        beneficiary_record=record, status=FisherfolkBeneficiaryConfirmationRequest.Status.PENDING,
                    )
                )
                if existing:
                    return Response(
                        {'detail': 'A confirmation request for this record is already pending admin review.'},
                        status=400,
                    )
                req = FisherfolkBeneficiaryConfirmationRequest.objects.create(
                    beneficiary_record=record,
                    submitted_by_organization_id=request.user.linked_record_id,
                    submitted_by=request.user,
                )
                for f in photos:
                    row = FisherfolkBeneficiaryConfirmationRequestPhoto.objects.create(confirmation_request=req, image=f, uploaded_by=request.user)
                    written.append(row.image.name)
        except Exception:
            discard_uploaded_files(written)
            raise

        notify_staff(
            actor=request.user, verb='submitted an on-behalf-of confirmation for',
            record_type='ff_beneficiary_obo', record_id=req.id,
            record_label=_beneficiary_label(record, record.fisherfolk),
        )
        return Response(FisherfolkBeneficiaryConfirmationRequestSerializer(req, context={'request': request}).data, status=drf_status.HTTP_201_CREATED)


class BeneficiaryConfirmationRequestListView(generics.ListAPIView):
    """Staff review queue — mirrors core.profile_self_service.StaffChangeRequestListView exactly."""
    serializer_class = BeneficiaryConfirmationRequestSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    pagination_class = None

    def get_queryset(self):
        qs = BeneficiaryConfirmationRequest.objects.select_related(
            'beneficiary_record__farmer', 'submitted_by_organization', 'submitted_by', 'reviewed_by',
        ).prefetch_related('photos')
        status_filter = self.request.query_params.get('status', 'PENDING').upper()
        if status_filter != 'ALL':
            qs = qs.filter(status=status_filter)
        return qs


class FisherfolkBeneficiaryConfirmationRequestListView(generics.ListAPIView):
    """Fisherfolk equivalent of BeneficiaryConfirmationRequestListView above."""
    serializer_class = FisherfolkBeneficiaryConfirmationRequestSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    pagination_class = None

    def get_queryset(self):
        qs = FisherfolkBeneficiaryConfirmationRequest.objects.select_related(
            'beneficiary_record__fisherfolk', 'submitted_by_organization', 'submitted_by', 'reviewed_by',
        ).prefetch_related('photos')
        status_filter = self.request.query_params.get('status', 'PENDING').upper()
        if status_filter != 'ALL':
            qs = qs.filter(status=status_filter)
        return qs


class BeneficiaryConfirmationRequestDecisionView(APIView):
    """
    POST .../<pk>/approve/  -> writes CONFIRMED onto the real BeneficiaryRecord,
                                in one atomic transaction with the request's
                                own status flip (mirrors
                                core.profile_self_service.StaffChangeRequestDecisionView).
    POST .../<pk>/reject/   -> leaves the real record's status untouched
                                (still AWAITING_CONFIRMATION) — the real
                                recipient, or a corrected resubmission, can
                                still legitimately confirm it later.
    """
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def post(self, request, pk, action):
        req = BeneficiaryConfirmationRequest.objects.select_related(
            'beneficiary_record__farmer', 'submitted_by_organization',
        ).filter(pk=pk).first()
        if req is None:
            return Response({'detail': 'Not found.'}, status=404)
        if req.status != BeneficiaryConfirmationRequest.Status.PENDING:
            return Response({'detail': f'This request is already {req.status.lower()}.'}, status=400)

        record = req.beneficiary_record
        if action == 'approve':
            if record.status != DistributionStatus.AWAITING_CONFIRMATION:
                return Response(
                    {'detail': (
                        f'The underlying record is no longer awaiting confirmation '
                        f'(current status: {record.get_status_display()}) — it may already '
                        f'have been resolved another way. This request cannot be approved as-is.'
                    )},
                    status=400,
                )
            with transaction.atomic():
                record.status = DistributionStatus.CONFIRMED
                record.confirmed_at = timezone.now()
                record.save()
                req.status = BeneficiaryConfirmationRequest.Status.APPROVED
                req.reviewed_by = request.user
                req.reviewed_at = timezone.now()
                req.save()
            verb, detail = 'approved your on-behalf-of confirmation for', ''
        elif action == 'reject':
            reason = _required_rejection_note(request)
            req.status = BeneficiaryConfirmationRequest.Status.REJECTED
            req.reviewed_by = request.user
            req.reviewed_at = timezone.now()
            req.rejection_reason = reason
            req.save()
            verb, detail = 'rejected your on-behalf-of confirmation for', reason
        else:
            return Response({'detail': f'Unknown action "{action}".'}, status=400)

        notify_user(
            recipient=req.submitted_by, actor=request.user, verb=verb,
            record_type='beneficiary_obo', record_id=req.id,
            record_label=_beneficiary_label(record, record.farmer), detail=detail,
        )
        return Response(BeneficiaryConfirmationRequestSerializer(req).data)


class FisherfolkBeneficiaryConfirmationRequestDecisionView(APIView):
    """Fisherfolk equivalent of BeneficiaryConfirmationRequestDecisionView above."""
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def post(self, request, pk, action):
        req = FisherfolkBeneficiaryConfirmationRequest.objects.select_related(
            'beneficiary_record__fisherfolk', 'submitted_by_organization',
        ).filter(pk=pk).first()
        if req is None:
            return Response({'detail': 'Not found.'}, status=404)
        if req.status != FisherfolkBeneficiaryConfirmationRequest.Status.PENDING:
            return Response({'detail': f'This request is already {req.status.lower()}.'}, status=400)

        record = req.beneficiary_record
        if action == 'approve':
            if record.status != DistributionStatus.AWAITING_CONFIRMATION:
                return Response(
                    {'detail': (
                        f'The underlying record is no longer awaiting confirmation '
                        f'(current status: {record.get_status_display()}) — it may already '
                        f'have been resolved another way. This request cannot be approved as-is.'
                    )},
                    status=400,
                )
            with transaction.atomic():
                record.status = DistributionStatus.CONFIRMED
                record.confirmed_at = timezone.now()
                record.save()
                req.status = FisherfolkBeneficiaryConfirmationRequest.Status.APPROVED
                req.reviewed_by = request.user
                req.reviewed_at = timezone.now()
                req.save()
            verb, detail = 'approved your on-behalf-of confirmation for', ''
        elif action == 'reject':
            reason = _required_rejection_note(request)
            req.status = FisherfolkBeneficiaryConfirmationRequest.Status.REJECTED
            req.reviewed_by = request.user
            req.reviewed_at = timezone.now()
            req.rejection_reason = reason
            req.save()
            verb, detail = 'rejected your on-behalf-of confirmation for', reason
        else:
            return Response({'detail': f'Unknown action "{action}".'}, status=400)

        notify_user(
            recipient=req.submitted_by, actor=request.user, verb=verb,
            record_type='ff_beneficiary_obo', record_id=req.id,
            record_label=_beneficiary_label(record, record.fisherfolk), detail=detail,
        )
        return Response(FisherfolkBeneficiaryConfirmationRequestSerializer(req).data)


# ─── Assistance Excel import ─────────────────────────────────────────────────
# "Import File" on the Crop/Fisheries Assistance pages — logs many
# assistance-given events from a spreadsheet instead of one at a time
# through AddAssistanceModal/LogGovernmentAssistanceModal. Shares the same
# dry_run convention as associations._run_organization_distribution and the
# same two-pass in-file-then-database duplicate detection as Crop/Fisheries
# Forecasting's imports (production/views.py ForecastingImportView,
# fisheries/views.py StockingForecastingImportView, associations/views.py
# _run_organization_import) — see any of those for the full rationale.
#
# Identity is (person, assistance_type, date_given): the three facts that
# together answer "is this the same assistance-giving event" from what this
# import captures. quantity/source_agency/remarks are details of the event,
# not part of what makes two rows "the same event" — same reasoning as
# excluding number_of_fingerlings from StockingRecord's identity.

ASSISTANCE_TEMPLATE_HEADERS = [
    'RSBSA Number', 'Assistance Type', 'Source Agency', 'Quantity',
    'Date Given (YYYY-MM-DD)', 'Season', 'Description / Remarks',
]
ASSISTANCE_TEMPLATE_COL_WIDTHS = [28, 20, 24, 18, 22, 12, 40]


CROP_ASSISTANCE_TEMPLATE_SAMPLE = [
    '01-33-02-014-058212', 'Seeds', 'Office of the Municipal Agriculturist', '5 sacks',
    '2026-01-15', 'Dry', 'Certified rice seeds distribution',
]
FISHERIES_ASSISTANCE_TEMPLATE_SAMPLE = [
    '01-33-02-999-000001', 'Fingerlings/Seeds', 'BFAR', '500 pcs',
    '2026-01-15', 'Dry', 'Bangus fingerlings distribution',
]
ASSISTANCE_ID_LABELS = {'farmer': 'Farmer RSBSA Number', 'fisherfolk': 'Fisherfolk RSBSA Number'}
ASSISTANCE_SAMPLES = {'farmer': CROP_ASSISTANCE_TEMPLATE_SAMPLE, 'fisherfolk': FISHERIES_ASSISTANCE_TEMPLATE_SAMPLE}
ASSISTANCE_NOTE_PREFIX = 'Assistance Type must be one of'


def _assistance_import_template_response(filename, id_label, sample_row, type_note):
    from io import BytesIO

    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = 'Assistance Import'

    headers = list(ASSISTANCE_TEMPLATE_HEADERS)
    headers[0] = id_label
    green_fill = PatternFill(start_color='1E4620', end_color='1E4620', fill_type='solid')
    white_font = Font(color='FFFFFF', bold=True)
    for col, (header, width) in enumerate(zip(headers, ASSISTANCE_TEMPLATE_COL_WIDTHS), 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = green_fill
        cell.font = white_font
        ws.column_dimensions[cell.column_letter].width = width

    for col, value in enumerate(sample_row, 1):
        ws.cell(row=2, column=col, value=value)
    ws.cell(row=4, column=1, value=type_note)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    response = HttpResponse(
        buf.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def _run_assistance_import(person_model, record_model, serializer_class, person_field, text_field, file, dry_run, user):
    """
    person_model/record_model: e.g. Farmer/BeneficiaryRecord.
    person_field: 'farmer' or 'fisherfolk' — both the serializer's FK field
    name and the record_model's own field name for it.
    text_field: 'remarks' (BeneficiaryRecord) or 'description'
    (FisherfolkBeneficiaryRecord) — mirrors OrganizationDistributionRequestSerializer's
    same per-model mapping in associations/views.py.
    """
    import datetime as dt
    from collections import defaultdict

    from core.utils import ImportFileError, cell_text, read_template_rows

    # The header row is located by content and must carry THIS page's id
    # label ("Farmer RSBSA Number" vs "Fisherfolk RSBSA Number"): a file from
    # the other page's template is refused up front instead of being half-read
    # (a person in both registries would otherwise get assistance recorded
    # under the wrong side).
    headers = list(ASSISTANCE_TEMPLATE_HEADERS)
    headers[0] = ASSISTANCE_ID_LABELS[person_field]
    try:
        rows_iter = read_template_rows(
            file, file.name, headers,
            sample_row=ASSISTANCE_SAMPLES[person_field], note_prefixes=(ASSISTANCE_NOTE_PREFIX,),
        )
    except ImportFileError as exc:
        return Response({'error': str(exc)}, status=400)
    except Exception as exc:
        return Response({'error': f'Invalid file: {exc}'}, status=400)

    valid_types = {label.upper(): value for value, label in record_model.AssistanceType.choices}

    errors = []
    rows_out = []
    ready = []
    groups = defaultdict(list)

    def invalid(row_num, msg):
        errors.append({'row': row_num, 'error': msg})
        rows_out.append({'row': row_num, 'status': 'invalid', 'label': f'Row {row_num}', 'detail': msg})

    def s(row, i):
        return cell_text(row[i]) if i < len(row) else ''

    for row_num, row in rows_iter:
        id_number = s(row, 0)
        type_raw = s(row, 1)
        source_agency = s(row, 2)
        quantity = s(row, 3)
        date_str = s(row, 4)
        season = s(row, 5)
        text_value = s(row, 6)
        if season.lower() in ('dry', 'wet'):
            season = season.title()  # "dry"/"WET" would otherwise fail the Dry/Wet choice check

        if not id_number or not type_raw or not date_str:
            invalid(row_num, 'RSBSA Number, Assistance Type, and Date Given are required.')
            continue

        person = person_model.objects.filter(rsbsa_number__iexact=id_number).first()
        if not person:
            invalid(row_num, f'No record found with RSBSA number "{id_number}".')
            continue

        assistance_type = valid_types.get(type_raw.upper())
        if not assistance_type:
            valid_list = ', '.join(v for v, _ in record_model.AssistanceType.choices)
            invalid(row_num, f'Invalid Assistance Type "{type_raw}". Use one of: {valid_list}.')
            continue

        try:
            date_given = dt.date.fromisoformat(date_str[:10])
        except ValueError:
            invalid(row_num, f'Invalid date "{date_str}". Use YYYY-MM-DD.')
            continue

        key = (person.id, assistance_type, date_given)
        ready.append({
            'row_num': row_num, 'person': person, 'assistance_type': assistance_type,
            'source_agency': source_agency, 'quantity': quantity, 'date_given': date_given,
            'season': season, 'text_value': text_value, 'key': key,
        })
        groups[key].append(row_num)

    created = 0
    for item in ready:
        row_num, key = item['row_num'], item['key']
        person, assistance_type, date_given = item['person'], item['assistance_type'], item['date_given']
        label = f'{person.full_name} — {assistance_type} ({date_given.isoformat()})'

        colliding_rows = [r for r in groups[key] if r != row_num]
        if colliding_rows:
            detail = (
                f'Duplicate of row(s) {", ".join(str(r) for r in colliding_rows)} in this file '
                f'(same person, Assistance Type "{assistance_type}", Date Given {date_given.isoformat()}). '
                f'None of the colliding rows were created — resolve the conflict and re-upload.'
            )
            rows_out.append({'row': row_num, 'status': 'duplicate_in_file', 'label': label, 'detail': detail, 'colliding_rows': colliding_rows})
            continue

        existing = record_model.objects.filter(
            **{person_field: person, 'assistance_type': assistance_type, 'date_given': date_given},
        ).first()
        if existing:
            detail = (
                f'A record already exists for {person.full_name}, Assistance Type "{assistance_type}", '
                f'Date Given {date_given.isoformat()} (record #{existing.id}). Not created again.'
            )
            rows_out.append({'row': row_num, 'status': 'duplicate_in_database', 'label': label, 'detail': detail, 'existing_record_id': existing.id})
            continue

        payload = {
            person_field: person.id,
            'assistance_type': assistance_type,
            # Import file still has one Source Agency column — wrapped into
            # a one-item list to match source_agencies' new shape, same
            # single-value-to-list conversion the data migration applies to
            # existing rows.
            'source_agencies': [item['source_agency']] if item['source_agency'] else [],
            'quantity': item['quantity'],
            'date_given': date_given,
            'season': item['season'],
            text_field: item['text_value'],
        }
        serializer = serializer_class(data=payload)
        if not serializer.is_valid():
            detail = '; '.join(f'{field}: {"; ".join(str(m) for m in msgs)}' for field, msgs in serializer.errors.items())
            invalid(row_num, detail or 'Invalid row.')
            continue

        if dry_run:
            rows_out.append({'row': row_num, 'status': 'ready', 'label': label, 'detail': 'Will be created.'})
            continue

        record = serializer.save(recorded_by=user)
        if record.status == DistributionStatus.AWAITING_CONFIRMATION:
            # An import loads assistance that already happened, so past-dated rows stay Confirmed (they must not
            # flood every recipient with old confirmations); future-dated rows stay Pending as before.
            record_model.objects.filter(pk=record.pk).update(status=DistributionStatus.CONFIRMED)
        created += 1
        rows_out.append({'row': row_num, 'status': 'ready', 'label': label, 'detail': f'Created (record #{record.id}).'})

    summary = {
        'ready': sum(1 for r in rows_out if r['status'] == 'ready'),
        'duplicate_in_file': sum(1 for r in rows_out if r['status'] == 'duplicate_in_file'),
        'duplicate_in_database': sum(1 for r in rows_out if r['status'] == 'duplicate_in_database'),
        'invalid': sum(1 for r in rows_out if r['status'] == 'invalid'),
    }
    return Response({
        'created': created, 'errors': errors, 'dry_run': dry_run, 'summary': summary, 'rows': rows_out,
    })


class CropAssistanceImportTemplateView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def get(self, request):
        return _assistance_import_template_response(
            'crop_assistance_import_template.xlsx',
            'Farmer RSBSA Number',
            CROP_ASSISTANCE_TEMPLATE_SAMPLE,
            'Assistance Type must be one of: Fertilizer / Seeds / Equipment / Cash Assistance / Training / Other',
        )


class CropAssistanceImportView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    parser_classes = [MultiPartParser]

    def post(self, request):
        file = request.FILES.get('file')
        if not file:
            return Response({'error': 'No file uploaded.'}, status=400)
        dry_run = str(request.data.get('dry_run', '')).strip().lower() in ('1', 'true', 'yes')
        from farmers.models import Farmer
        from core.utils import run_import_atomically
        return run_import_atomically(lambda: _run_assistance_import(
            Farmer, BeneficiaryRecord, BeneficiaryRecordCreateSerializer,
            'farmer', 'remarks', file, dry_run, request.user,
        ))


class FisheriesAssistanceImportTemplateView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def get(self, request):
        return _assistance_import_template_response(
            'fisheries_assistance_import_template.xlsx',
            'Fisherfolk RSBSA Number',
            FISHERIES_ASSISTANCE_TEMPLATE_SAMPLE,
            'Assistance Type must be one of: Fishing Gear / Boat/Vessel / Fingerlings/Seeds / Feed / '
            'Equipment / Cash Assistance / Training / Other',
        )


class FisheriesAssistanceImportView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    parser_classes = [MultiPartParser]

    def post(self, request):
        file = request.FILES.get('file')
        if not file:
            return Response({'error': 'No file uploaded.'}, status=400)
        dry_run = str(request.data.get('dry_run', '')).strip().lower() in ('1', 'true', 'yes')
        from fisheries.models import Fisherfolk
        from core.utils import run_import_atomically
        return run_import_atomically(lambda: _run_assistance_import(
            Fisherfolk, FisherfolkBeneficiaryRecord, FisherfolkBeneficiaryRecordCreateSerializer,
            'fisherfolk', 'description', file, dry_run, request.user,
        ))
