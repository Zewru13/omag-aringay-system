"""
One small "document" model shared by the new Reports page reports.

A report builds a plain dict:

    {
      'title':    'Registered Farmers',
      'subtitle': 'Barangay: Alaska',                       # optional one line under the title
      'stats':    [['Total Farmers', 12], ...],             # optional summary figures
      'sections': [                                         # one or more
          {'title': 'Details', 'pairs': [['RSBSA No.', '01-...'], ...]},                      # label / value grid
          {'title': 'Farm Parcels', 'headers': ['Parcel', 'Area'], 'rows': [[...], ...],      # table
           'empty': 'No parcels on record.'},
      ],
    }

and this module turns the SAME dict into the page's JSON preview, the official PDF (letterhead, signature block) and the Excel
file, so the three can never disagree. Every value is a string or number; the JSON is exactly what the preview table shows.
"""
import io
from datetime import date

from django.http import HttpResponse
from reportlab.lib.pagesizes import landscape, letter
from reportlab.platypus import Spacer

from core.exports import SafeWorkbook
from rest_framework.response import Response

from .period import period_text
from .pdf_utils import (
    NumberedCanvas,
    build_keyvalue_grid,
    build_ledger_table,
    build_official_header,
    build_section_heading,
    build_signature_block,
    build_stats_row,
    fmt_field,
)
from .views import LOGO_PATH, PAGE_WIDTH, _empty_state, _finish_pdf, _make_doc, _sig_params, _today

ROMAN = ['I', 'II', 'III', 'IV', 'V', 'VI', 'VII', 'VIII', 'IX', 'X']


def make_document(title, subtitle='', stats=None, sections=None, heading=None, period='', detail=None):
    """`title` / `subtitle` are what the page preview and the Excel show. `heading` is the formal title printed on the PDF (the
    Print and Preview PDF output), `period` the reporting-period line under it (see reports.period.period_text); both are only
    used by the PDF and the preview, so the Excel file is untouched. `detail` is the PDF's line under the period: the filters, without
    the dates that the period line already says (left out, it falls back to the title and subtitle)."""
    return {'title': title, 'subtitle': subtitle or '', 'stats': stats or [], 'sections': sections or [],
            'heading': heading or title, 'period': period or '', 'detail': detail}


class ReportPageCanvas(NumberedCanvas):
    """Page footer of the Reports page PDFs: page numbers and a plain line, without the system-name tagline."""
    disclaimer = 'This is an official record of the Office of the Municipal Agriculturist.'


# The seal is drawn at 80 points wide, its height following its own proportions (same as the individual disaster reports).
LOGO_WIDTH = 80


def filters_line(doc):
    """The line under the period: what the report is and how it was filtered. Left out when it only repeats the heading."""
    if doc.get('detail') is not None:
        return doc['detail']
    parts = []
    if doc.get('title') and doc['title'] != doc.get('heading'):
        parts.append(doc['title'])
    if doc.get('subtitle'):
        parts.append(doc['subtitle'])
    return ' · '.join(parts)


def text(value):
    """A value as the report shows it: blank / None become a dash-free 'N/A'-like empty string handled by the caller."""
    if value is None:
        return ''
    if isinstance(value, date):
        return value.strftime('%b %d, %Y')
    return str(value)


WIDE_TABLE_COLUMNS = 8       # a table with this many columns gets a landscape page
LANDSCAPE_WIDTH = 648        # 792 - 72 - 72


def _longest_word(values):
    return max((len(word) for value in values for word in str(value).split()), default=1)


def column_widths(headers, rows, total=PAGE_WIDTH):
    """Column widths (points) that add up exactly to `total`. Each column first gets room for its longest single word (so
    "Female" or "Inactive" never breaks mid-word), then the rest is shared by how much text the column holds."""
    weights, minimums = [], []
    for i, header in enumerate(headers):
        column = [row[i] for row in rows if i < len(row)]
        longest = max([len(str(header)) * 0.7] + [min(len(str(v)), 34) for v in column])
        weights.append(max(longest, 6))
        minimums.append(_longest_word([header] + column[:200]) * 4.6 + 12)
    extra = total - sum(minimums)
    if extra <= 0:
        scale = total / sum(minimums)
        widths = [round(m * scale, 1) for m in minimums]
    else:
        spare = sum(weights)
        widths = [round(m + extra * w / spare, 1) for m, w in zip(minimums, weights)]
    widths[-1] = round(total - sum(widths[:-1]), 1)
    return widths


def _is_wide(doc):
    return any(len(section.get('headers', ())) >= WIDE_TABLE_COLUMNS for section in doc.get('sections', []))


def as_json(doc):
    return Response(doc)


def build_pdf(doc, request, filename_stem):
    sig = _sig_params(request)
    wide = _is_wide(doc)
    width = LANDSCAPE_WIDTH if wide else PAGE_WIDTH
    buf, pdf = _make_doc()
    # Slim top and bottom margins (the footer line sits in the bottom margin): the letterhead is tall, so the table starts on page 1.
    pdf.topMargin, pdf.bottomMargin = 40, 50
    if wide:
        pdf.pagesize = landscape(letter)
        pdf.width = width
    story = build_official_header(
        LOGO_PATH, (doc.get('heading') or doc['title']).upper(), doc.get('period') or None,
        logo_width=LOGO_WIDTH, detail=filters_line(doc) or None, compact=True,
    )
    if doc.get('stats'):
        story.append(build_stats_row([(label, value) for label, value in doc['stats']], page_width=width))
        story.append(Spacer(1, 10))
    sections = doc.get('sections', [])
    for index, section in enumerate(sections):
        if len(sections) > 1 or section.get('title'):
            story += build_section_heading(ROMAN[index] if index < len(ROMAN) else str(index + 1), section.get('title', ''))
        if 'pairs' in section:
            story.append(build_keyvalue_grid([(label, fmt_field(value)) for label, value in section['pairs']], columns=2, page_width=width))
        else:
            headers, rows = section['headers'], section['rows']
            if rows:
                story.append(build_ledger_table(headers, rows, col_widths=column_widths(headers, rows, width)))
            else:
                story.append(_empty_state(section.get('empty', 'No records found.')))
    story.append(Spacer(1, 30))
    story.append(build_signature_block(
        left_label='Prepared by',
        left_name=sig.get('prepared_by') or None,
        left_role=sig.get('position') or 'OMAg Records Officer',
        reviewed_by=sig.get('reviewed_by'),
        reviewed_role=sig.get('reviewed_role'),
        noted_by=sig.get('noted_by'),
        noted_role=sig.get('noted_role'),
        page_width=width,
    ))
    pdf.build(story, canvasmaker=ReportPageCanvas)
    return _finish_pdf(buf, f'{filename_stem}_{_today()}.pdf')


SHEET_TITLE_FORBIDDEN = r'\/?*[]:'


def _sheet_title(title):
    """A worksheet name Excel accepts: at most 31 characters and none of the characters  \\ / ? * [ ] :"""
    cleaned = ''.join(' ' if ch in SHEET_TITLE_FORBIDDEN else ch for ch in (title or '')).strip()
    return ' '.join(cleaned.split())[:31] or 'Report'


def build_excel(doc, filename_stem):
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = SafeWorkbook()
    ws = wb.active
    ws.title = _sheet_title(doc['title'])
    header_fill = PatternFill('solid', fgColor='1E4620')
    header_font = Font(bold=True, color='FFFFFF', size=10)

    ws.append(['Republic of the Philippines'])
    ws.append(['Office of the Municipal Agriculturist — Aringay, La Union'])
    ws.append([doc['title']])
    if doc.get('subtitle'):
        ws.append([doc['subtitle']])
    ws.append([f'Date Generated: {_today()}'])
    for label, value in doc.get('stats', []):
        ws.append([label, value])
    for section in doc.get('sections', []):
        ws.append([])
        if section.get('title'):
            ws.append([section['title']])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=11)
        if 'pairs' in section:
            for label, value in section['pairs']:
                ws.append([label, fmt_field(value)])
                ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        else:
            ws.append(list(section['headers']))
            for cell in ws[ws.max_row]:
                cell.fill, cell.font, cell.alignment = header_fill, header_font, Alignment(horizontal='center')
            if section['rows']:
                for row in section['rows']:
                    ws.append(list(row))
            else:
                ws.append([section.get('empty', 'No records found.')])
    for column in ws.columns:
        longest = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        ws.column_dimensions[column[0].column_letter].width = min(max(longest + 2, 10), 48)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    response = HttpResponse(out.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename_stem}_{_today()}.xlsx"'
    return response


def respond(doc, request, filename_stem):
    """The report in the format the request asks for: format=json (the page preview, default), pdf or excel."""
    fmt = request.query_params.get('format', 'json').lower()
    if fmt == 'pdf':
        return build_pdf(doc, request, filename_stem)
    if fmt in ('excel', 'xlsx'):
        return build_excel(doc, filename_stem)
    return as_json(doc)
