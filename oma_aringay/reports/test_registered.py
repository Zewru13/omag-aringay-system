"""Report 1 (Registered Farmers / Fisherfolk / Organizations): bulk with filters, one record, JSON + PDF + Excel."""
import io
from datetime import date

import openpyxl
from django.contrib.auth import get_user_model
from rest_framework.test import APITestCase

from associations.models import (
    FarmersOrganization, FarmersOrganizationMember, FisherfolkOrganization, FisherfolkOrganizationMember,
    OrganizationOfficer,
)
from farmers.models import FarmParcel, Farmer
from fisheries.models import Fisherfolk

URL = '/api/reports/registered/'


def _farmer(rsbsa, first, last, barangay='Alaska', **over):
    data = dict(rsbsa_number=rsbsa, first_name=first, last_name=last, gender='Female', birthdate=date(1985, 3, 3),
                barangay=barangay, civil_status='Single')
    data.update(over)
    return Farmer.objects.create(**data)


def _fisher(first, last, rsbsa, barangay='Alaska', **over):
    data = dict(rsbsa_number=rsbsa, first_name=first, last_name=last, gender='Male', birthdate=date(1980, 1, 1), barangay=barangay,
                fisherfolk_type='Aquaculture')
    data.update(over)
    return Fisherfolk.objects.create(**data)


def _org(model, name, number, barangay='Alaska', **over):
    data = dict(name=name, barangay=barangay, contact_person='Contact', contact_number='0917', registration_number=number,
                date_organized=date(2020, 1, 1), date_registered=date(2021, 2, 2))
    data.update(over)
    return model.objects.create(**data)


class RegisteredReportTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        User = get_user_model()
        cls.admin = User.objects.create_user(username='admin1', password='pw', role='ADMIN', email='a@x.test')
        cls.citizen = User.objects.create_user(username='cit1', password='pw', role='CITIZEN', email='c@x.test')
        cls.maria = _farmer('01-033-02-019-000151', 'Maria', 'Santos')
        cls.jose = _farmer('01-033-02-019-000152', 'Jose', 'Reyes', barangay='Baracbac')
        FarmParcel.objects.create(farmer=cls.maria, parcel_number='P-1', area_hectares='1.50', barangay='Alaska')
        cls.org = _org(FarmersOrganization, 'Alaska Farmers Org', 'ORG-1')
        FarmersOrganizationMember.objects.create(organization=cls.org, farmer=cls.maria, date_joined=date(2022, 1, 1))
        OrganizationOfficer.objects.create(organization=cls.org, farmer=cls.maria, position='President', full_name='Maria Santos')
        cls.other_org = _org(FarmersOrganization, 'Baracbac Org', 'ORG-2', barangay='Baracbac')
        cls.fisher = _fisher('Pedro', 'Cruz', '01-033-02-019-000201')
        cls.fisher2 = _fisher('Ana', 'Lim', '01-033-02-019-000202', barangay='Baracbac')
        cls.forg = _org(FisherfolkOrganization, 'Alaska Fisherfolk Org', 'FORG-1')
        FisherfolkOrganizationMember.objects.create(organization=cls.forg, fisherfolk=cls.fisher, date_joined=date(2022, 1, 1))

    def setUp(self):
        self.client.force_authenticate(self.admin)

    def get(self, **params):
        r = self.client.get(URL, params)
        self.assertEqual(r.status_code, 200, r.content[:300])
        return r.json()

    def rows(self, doc):
        return doc['sections'][0]['rows']

    # --- access and input checks
    def test_citizens_are_refused(self):
        self.client.force_authenticate(self.citizen)
        self.assertEqual(self.client.get(URL, {'kind': 'farmers'}).status_code, 403)

    def test_a_missing_or_unknown_kind_is_a_400(self):
        self.assertEqual(self.client.get(URL).status_code, 400)
        self.assertEqual(self.client.get(URL, {'kind': 'trees'}).status_code, 400)

    def test_bad_dates_and_range_are_a_400(self):
        self.assertEqual(self.client.get(URL, {'kind': 'farmers', 'registered_from': 'soon'}).status_code, 400)
        self.assertEqual(self.client.get(URL, {'kind': 'farmers', 'registered_from': '2030-01-02', 'registered_to': '2030-01-01'}).status_code, 400)
        self.assertEqual(self.client.get(URL, {'kind': 'farmers', 'id': 'abc'}).status_code, 400)

    # --- bulk
    def test_bulk_farmers_lists_everyone_with_their_parcel_area_and_organization(self):
        doc = self.get(kind='farmers')
        self.assertEqual(doc['title'], 'Registered Farmers')
        self.assertEqual(doc['stats'], [['Total Farmers', 2]])
        by_name = {r[1]: r for r in self.rows(doc)}
        maria = by_name['Maria Santos']
        self.assertEqual(maria[0], '01-033-02-019-000151')
        self.assertEqual(maria[5], 1)
        self.assertEqual(maria[6], '1.50')
        self.assertEqual(maria[7], 'Alaska Farmers Org')
        self.assertEqual(len(maria), 9)    # the last column is the date registered: there is no Status column
        self.assertNotIn('Active', maria)

    def test_bulk_filters_narrow_by_barangay_and_search(self):
        self.assertEqual([r[1] for r in self.rows(self.get(kind='farmers', barangay='Baracbac'))], ['Jose Reyes'])
        self.assertEqual([r[1] for r in self.rows(self.get(kind='farmers', search='000152'))], ['Jose Reyes'])
        self.assertEqual([r[1] for r in self.rows(self.get(kind='farmers', search='maria santos'))], ['Maria Santos'])

    def test_bulk_date_registered_range_is_inclusive_of_today_and_excludes_the_future(self):
        today = date.today().isoformat()
        self.assertEqual(len(self.rows(self.get(kind='farmers', registered_from=today, registered_to=today))), 2)
        self.assertEqual(len(self.rows(self.get(kind='farmers', registered_from='2999-01-01'))), 0)

    def test_an_empty_result_is_a_normal_document_with_its_empty_message(self):
        doc = self.get(kind='farmers', search='zzzz')
        self.assertEqual(self.rows(doc), [])
        self.assertEqual(doc['stats'], [['Total Farmers', 0]])
        self.assertTrue(doc['sections'][0]['empty'])

    def test_bulk_fisherfolk(self):
        doc = self.get(kind='fisherfolk')
        self.assertEqual(doc['title'], 'Registered Fisherfolk')
        self.assertEqual(sorted(r[1] for r in self.rows(doc)), ['Ana Lim', 'Pedro Cruz'])
        self.assertEqual([r[1] for r in self.rows(self.get(kind='fisherfolk', barangay='Alaska'))], ['Pedro Cruz'])
        pedro = next(r for r in self.rows(doc) if r[1] == 'Pedro Cruz')
        self.assertEqual(pedro[6], 'Alaska Fisherfolk Org')

    def test_bulk_organizations_of_both_kinds(self):
        doc = self.get(kind='farmers_orgs')
        self.assertEqual(sorted(r[0] for r in self.rows(doc)), ['Alaska Farmers Org', 'Baracbac Org'])
        alaska = next(r for r in self.rows(doc) if r[0] == 'Alaska Farmers Org')
        self.assertEqual(alaska[5], 1)  # active members
        self.assertEqual([r[0] for r in self.rows(self.get(kind='farmers_orgs', search='ORG-2'))], ['Baracbac Org'])
        self.assertEqual([r[0] for r in self.rows(self.get(kind='fisherfolk_orgs'))], ['Alaska Fisherfolk Org'])
        self.assertEqual(self.rows(self.get(kind='fisherfolk_orgs', registered_from='2022-01-01')), [])

    # --- one record
    def test_one_farmer_is_a_profile_with_details_parcels_and_organizations(self):
        doc = self.get(kind='farmers', id=self.maria.id)
        self.assertEqual(doc['title'], 'Farmer Registration Record')
        self.assertEqual(doc['subtitle'], 'Maria Santos')
        self.assertEqual([s['title'] for s in doc['sections']], ['Farmer Details', 'Farm Parcels', 'Organizations'])
        details = dict(doc['sections'][0]['pairs'])
        self.assertEqual(details['RSBSA No.'], '01-033-02-019-000151')
        self.assertEqual(doc['sections'][1]['rows'][0][0], 'P-1')
        self.assertEqual(doc['sections'][2]['rows'][0][0], 'Alaska Farmers Org')

    def test_one_fisherfolk_and_one_organization(self):
        doc = self.get(kind='fisherfolk', id=self.fisher.id)
        self.assertEqual(doc['subtitle'], 'Pedro Cruz')
        self.assertEqual(doc['sections'][1]['rows'][0][0], 'Alaska Fisherfolk Org')
        org = self.get(kind='farmers_orgs', id=self.org.id)
        self.assertEqual([s['title'] for s in org['sections']], ['Organization Details', 'Officers', 'Members'])
        self.assertEqual(org['sections'][1]['rows'][0][:2], ['President', 'Maria Santos'])
        self.assertEqual(org['sections'][2]['rows'][0][0], 'Maria Santos')
        forg = self.get(kind='fisherfolk_orgs', id=self.forg.id)
        self.assertEqual(forg['sections'][2]['rows'][0][0], 'Pedro Cruz')

    def test_a_record_id_that_does_not_exist_is_a_404(self):
        from django.http import Http404
        from .registered import one_farmer
        with self.assertRaises(Http404):
            one_farmer(999999)

    def test_one_record_ignores_the_bulk_filters(self):
        doc = self.get(kind='farmers', id=self.maria.id, barangay='Baracbac')
        self.assertEqual(doc['subtitle'], 'Maria Santos')

    # --- the "one record" picker
    def test_options_lists_matching_records_for_the_picker(self):
        r = self.client.get(URL, {'kind': 'farmers', 'options': 1, 'search': 'maria'})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json(), [{'id': self.maria.id, 'label': 'Maria Santos', 'detail': '01-033-02-019-000151 · Alaska'}])
        orgs = self.client.get(URL, {'kind': 'fisherfolk_orgs', 'options': 1}).json()
        self.assertEqual([o['label'] for o in orgs], ['Alaska Fisherfolk Org'])
        self.assertEqual(self.client.get(URL, {'kind': 'trees', 'options': 1}).status_code, 400)

    def test_options_are_capped(self):
        for i in range(30):
            _farmer(f'01-033-02-019-{i + 500:06d}', f'Extra{i}', 'Many')
        self.assertEqual(len(self.client.get(URL, {'kind': 'farmers', 'options': 1}).json()), 25)

    # --- the three formats come from the same document
    def test_pdf_is_a_pdf(self):
        r = self.client.get(URL, {'kind': 'farmers', 'format': 'pdf', 'prepared_by': 'Records Clerk'})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r['Content-Type'], 'application/pdf')
        self.assertTrue(r.content.startswith(b'%PDF'))
        one = self.client.get(URL, {'kind': 'farmers_orgs', 'id': self.org.id, 'format': 'pdf'})
        self.assertTrue(one.content.startswith(b'%PDF'))

    def test_a_table_with_many_columns_is_printed_on_a_landscape_page_and_a_profile_stays_portrait(self):
        wide = self.client.get(URL, {'kind': 'farmers', 'format': 'pdf'})
        self.assertIn(b'/MediaBox [ 0 0 792 612 ]', wide.content)
        profile = self.client.get(URL, {'kind': 'farmers', 'id': self.maria.id, 'format': 'pdf'})
        self.assertIn(b'/MediaBox [ 0 0 612 792 ]', profile.content)

    def test_column_widths_fill_the_page_and_keep_short_words_whole(self):
        from .documents import column_widths
        headers = ['RSBSA No.', 'Name', 'Sex', 'Status']
        rows = [['01-033-02-019-000151', 'Maria Santos Dela Cruz', 'Female', 'Inactive']]
        for total in (468, 648):
            widths = column_widths(headers, rows, total)
            self.assertAlmostEqual(sum(widths), total, places=1)
            self.assertGreaterEqual(widths[2], len('Female') * 4.6 + 12 - 0.1)
            self.assertGreaterEqual(widths[3], len('Inactive') * 4.6 + 12 - 0.1)

    def test_an_empty_pdf_still_builds(self):
        r = self.client.get(URL, {'kind': 'fisherfolk', 'search': 'zzzz', 'format': 'pdf'})
        self.assertTrue(r.content.startswith(b'%PDF'))

    def test_excel_has_the_same_rows_as_the_preview(self):
        r = self.client.get(URL, {'kind': 'farmers', 'format': 'excel'})
        self.assertEqual(r.status_code, 200)
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        cells = [str(c) for row in ws.iter_rows(values_only=True) for c in row if c is not None]
        self.assertIn('Maria Santos', cells)
        self.assertIn('Jose Reyes', cells)
        self.assertIn('Total Farmers', cells)

    def test_excel_text_that_looks_like_a_formula_is_neutralized(self):
        _farmer('01-033-02-019-000999', '=1+1', 'Evil')
        r = self.client.get(URL, {'kind': 'farmers', 'format': 'excel', 'search': 'Evil'})
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        cells = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and 'Evil' in c.value]
        self.assertTrue(cells)
        self.assertTrue(all(c.data_type == 's' for c in cells))
