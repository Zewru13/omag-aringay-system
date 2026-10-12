"""Importing an organization's equipment from a file (both domains)."""
import io
from datetime import date, timedelta
from unittest import mock

from django.contrib.auth import get_user_model
from django.db import IntegrityError
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from associations.test_member_import import DOMAINS, make_citizen, make_org, make_person, sheet

HEADERS = ['Equipment', 'Quantity', 'Condition', 'Acquisition Date (YYYY-MM-DD)', 'Remarks']
KNOWN = {'farmers': 'Hand Tractor', 'fisherfolk': 'Fish Cage'}
TODAY = date.today().isoformat()


def equipment_model(domain):
    from associations import org_imports
    return org_imports.DOMAINS[domain]['equipment_model']


class EquipmentImportTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        cls.staff = get_user_model().objects.create_user(
            username='ei-staff', password='pw', role='ADMIN', email='ei-staff@x.test')

    def setUp(self):
        self.client.force_authenticate(self.staff)

    def url(self, domain, org, suffix='equipment/import/'):
        return f'/api/organizations/{domain}/{org.id}/{suffix}'

    def run_import(self, domain, org, rows, dry_run=False, headers=HEADERS):
        return self.client.post(self.url(domain, org), {'file': sheet(rows, headers, 'equipment.xlsx'),
                                                        'dry_run': 'true' if dry_run else 'false'}, format='multipart')

    def statuses(self, response):
        self.assertEqual(response.status_code, 200, response.content)
        return [row['status'] for row in response.data['rows']]

    def test_template_has_the_headers_and_lists_the_valid_equipment(self):
        for domain in DOMAINS:
            response = self.client.get(self.url(domain, make_org(domain), 'equipment/import/template/'))
            self.assertEqual(response.status_code, 200)
            ws = load_workbook(io.BytesIO(response.content)).active
            self.assertEqual([c.value for c in ws[1]][:5], HEADERS)
            notes = [str(c.value) for row in ws.iter_rows() for c in row if str(c.value or '').startswith('Valid equipment:')]
            self.assertEqual(len(notes), 1)
            self.assertIn(KNOWN[domain], notes[0])

    def test_known_names_match_ignoring_case_and_spacing(self):
        for domain in DOMAINS:
            org = make_org(domain)
            name = KNOWN[domain].lower().replace(' ', '  ')
            response = self.run_import(domain, org, [[f' {name} ', 2, 'good', TODAY, 'fine']])
            self.assertEqual(self.statuses(response), ['ready'], domain)
            item = equipment_model(domain).objects.get(organization=org)
            self.assertEqual((item.equipment_name, item.quantity, item.condition, item.remarks), (KNOWN[domain], 2, 'Good', 'fine'))

    def test_unknown_name_is_recorded_as_other_with_the_text(self):
        for domain in DOMAINS:
            org = make_org(domain)
            response = self.run_import(domain, org, [['Water Buffalo', 1, 'Good', '', '']])
            self.assertEqual(self.statuses(response), ['ready'])
            self.assertEqual(response.data['rows'][0]['detail'].count('Recorded as Other: Water Buffalo'), 1)
            item = equipment_model(domain).objects.get(organization=org)
            self.assertEqual((item.equipment_name, item.equipment_name_other), ('Other', 'Water Buffalo'))

    def test_other_without_text_and_blank_names_are_invalid(self):
        for domain in DOMAINS:
            response = self.run_import(domain, make_org(domain), [['Other', 1, 'Good', '', ''], ['', 1, 'Good', '', '']])
            self.assertEqual(self.statuses(response), ['invalid', 'invalid'])

    def test_quantity_rules(self):
        for domain in DOMAINS:
            org = make_org(domain)
            name = KNOWN[domain]
            rows = [[name, 2.0, 'Good', '', ''], [name, 0, 'Good', '', ''], [name, 1.5, 'Good', '', ''], [name, 'abc', 'Good', '', ''],
                    [name, '', 'Good', '', '']]
            response = self.run_import(domain, org, rows)
            self.assertEqual(self.statuses(response), ['ready', 'invalid', 'invalid', 'invalid', 'invalid'], domain)
            self.assertEqual(equipment_model(domain).objects.get(organization=org).quantity, 2)

    def test_condition_and_date_rules(self):
        for domain in DOMAINS:
            name = KNOWN[domain]
            future = (date.today() + timedelta(days=5)).isoformat()
            rows = [[name, 1, 'Shiny', '', ''], [name, 1, 'Needs Repair', future, ''], [name, 1, 'Not Operational', 'junk', ''],
                    [name, 1, 'Needs Repair', '', '']]
            response = self.run_import(domain, make_org(domain), rows)
            self.assertEqual(self.statuses(response), ['invalid', 'invalid', 'invalid', 'ready'], domain)

    def test_identical_stored_row_is_already_exists_and_a_repeat_in_the_file_is_a_duplicate(self):
        for domain in DOMAINS:
            org = make_org(domain)
            name = KNOWN[domain]
            equipment_model(domain).objects.create(
                organization=org, equipment_name=name, quantity=2, condition='Good', acquisition_date=date(2025, 3, 1))
            rows = [[name, 2, 'Good', '2025-03-01', ''], [name, 3, 'Good', '2025-03-01', ''], [name, 3, 'good', '2025-03-01', 'again']]
            response = self.run_import(domain, org, rows)
            self.assertEqual(self.statuses(response), ['already_exists', 'ready', 'duplicate_in_file'], domain)
            self.assertEqual(equipment_model(domain).objects.filter(organization=org).count(), 2)

    def test_dry_run_changes_nothing_and_predicts_the_confirm(self):
        for domain in DOMAINS:
            org = make_org(domain)
            name = KNOWN[domain]
            rows = [[name, 2, 'Good', '', ''], [name, 2, 'Good', '', ''], ['Other', 1, 'Good', '', '']]
            preview = self.run_import(domain, org, rows, dry_run=True)
            self.assertEqual(equipment_model(domain).objects.filter(organization=org).count(), 0)
            confirm = self.run_import(domain, org, rows)
            self.assertEqual(self.statuses(preview), self.statuses(confirm))
            self.assertEqual(self.statuses(confirm), ['ready', 'duplicate_in_file', 'invalid'])
            self.assertEqual(equipment_model(domain).objects.filter(organization=org).count(), 1)

    def test_an_unexpected_database_error_saves_nothing(self):
        from associations import org_imports
        for domain in DOMAINS:
            serializer = org_imports.DOMAINS[domain]['equipment_serializer']
            real_save = serializer.save
            calls = {'n': 0}

            def second_fails(instance, *args, **kwargs):
                calls['n'] += 1
                if calls['n'] == 2:
                    raise IntegrityError('simulated')
                return real_save(instance, *args, **kwargs)

            org = make_org(domain)
            rows = [[KNOWN[domain], 1, 'Good', '', ''], [KNOWN[domain], 2, 'Good', '', '']]
            with mock.patch.object(serializer, 'save', second_fails):
                response = self.run_import(domain, org, rows)
            self.assertEqual(response.status_code, 409, domain)
            self.assertEqual(equipment_model(domain).objects.filter(organization=org).count(), 0)

    def test_wrong_header_and_too_many_rows_are_refused(self):
        for domain in DOMAINS:
            org = make_org(domain)
            self.assertEqual(self.run_import(domain, org, [['x']], headers=['Name', 'Barangay']).status_code, 400)
            rows = [[KNOWN[domain], 1, 'Good', '', ''] for _ in range(1001)]
            self.assertEqual(self.run_import(domain, org, rows).status_code, 400)

    def test_permissions(self):
        for domain, cfg in DOMAINS.items():
            org, other = make_org(domain), make_org(domain)
            rows = [[KNOWN[domain], 1, 'Good', '', '']]
            own = make_citizen(f'ei-own-{domain}', cfg['link'], org.id)
            self.client.force_authenticate(own)
            self.assertEqual(self.run_import(domain, org, rows).status_code, 200, domain)
            self.assertIn(self.run_import(domain, other, rows).status_code, (403, 404), domain)
            other_domain = 'fisherfolk' if domain == 'farmers' else 'farmers'
            self.assertIn(self.run_import(other_domain, make_org(other_domain), rows).status_code, (403, 404), domain)
            person = make_person(domain)
            person_user = make_citizen(f'ei-person-{domain}', 'FARMER' if domain == 'farmers' else 'FISHERFOLK', person.id)
            self.client.force_authenticate(person_user)
            self.assertIn(self.run_import(domain, org, rows).status_code, (403, 404), domain)
            self.client.force_authenticate(None)
            self.assertEqual(self.run_import(domain, org, rows).status_code, 401, domain)
            self.client.force_authenticate(self.staff)
