"""File imports of an organization's members and equipment (both domains).

Members are people who are ALREADY registered, matched by RSBSA number, and are created with exactly the logic of the
Add Member view (the member serializer, then `_resolve_and_assign_main`). Every row runs in its own savepoint and a dry
run rolls the whole import back, so the preview shows what the confirm will do, including the main-organization rule.
"""
import datetime
from io import BytesIO

from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from rest_framework.parsers import MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from accounts.permissions import IsOwnOrgOrStaff
from core.idnumbers import by_id_key
from core.import_batches import record_created, start_batch
from core.utils import ImportFileError, cell_text, is_plain_int, read_template_rows, run_import_atomically
from farmers.models import Farmer
from fisheries.models import Fisherfolk

from .models import (
    FarmersOrganization, FarmersOrganizationEquipment, FarmersOrganizationMember,
    FisherfolkOrganization, FisherfolkOrganizationEquipment, FisherfolkOrganizationMember,
)
from .serializers import (
    FarmersOrganizationEquipmentSerializer, FarmersOrganizationMemberSerializer,
    FisherfolkOrganizationEquipmentSerializer, FisherfolkOrganizationMemberSerializer,
)
from .views import MainChoiceRequired, _resolve_and_assign_main
from core.exports import SafeWorkbook

DOMAINS = {
    'farmers': {
        'link_type': 'FARMERS_ORG',
        'org_model': FarmersOrganization,
        'person_model': Farmer, 'person_fk': 'farmer',
        'member_model': FarmersOrganizationMember, 'member_serializer': FarmersOrganizationMemberSerializer,
        'equipment_model': FarmersOrganizationEquipment, 'equipment_serializer': FarmersOrganizationEquipmentSerializer,
    },
    'fisherfolk': {
        'link_type': 'FISHERFOLK_ORG',
        'org_model': FisherfolkOrganization,
        'person_model': Fisherfolk, 'person_fk': 'fisherfolk',
        'member_model': FisherfolkOrganizationMember, 'member_serializer': FisherfolkOrganizationMemberSerializer,
        'equipment_model': FisherfolkOrganizationEquipment, 'equipment_serializer': FisherfolkOrganizationEquipmentSerializer,
    },
}

MEMBER_HEADERS = ['RSBSA Number', 'Position', 'Date Joined (YYYY-MM-DD)']
MEMBER_WIDTHS = [26, 22, 26]
MEMBER_NOTE_PREFIX = 'Valid positions:'
MEMBER_STATUSES = ('ready', 'already_member', 'not_found', 'needs_main_choice', 'duplicate_in_file', 'invalid')


def parse_import_date(raw):
    """A date cell as a date: blank gives None; an Excel date/datetime cell or YYYY-MM-DD text is accepted; anything
    else raises ValueError."""
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return None
    if isinstance(raw, datetime.datetime):
        return raw.date()
    if isinstance(raw, datetime.date):
        return raw
    try:
        return datetime.date.fromisoformat(str(raw).strip()[:10])
    except ValueError:
        raise ValueError(f'Invalid date "{raw}". Use YYYY-MM-DD.')


def _template_response(filename, headers, widths, sample_row, note):
    from openpyxl.styles import Font, PatternFill

    book = SafeWorkbook()
    ws = book.active
    ws.title = 'Import'
    green_fill = PatternFill(start_color='1E4620', end_color='1E4620', fill_type='solid')
    white_font = Font(color='FFFFFF', bold=True)
    for col, (header, width) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = green_fill
        cell.font = white_font
        ws.column_dimensions[cell.column_letter].width = width
    for col, value in enumerate(sample_row, 1):
        ws.cell(row=2, column=col, value=value)
    ws.cell(row=4, column=1, value=note)
    buffer = BytesIO()
    book.save(buffer)
    response = HttpResponse(
        buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def _read_rows(file, headers, sample_row, note_prefix):
    """The template's data rows, or an error Response."""
    try:
        return read_template_rows(file, file.name, headers, sample_row=sample_row, note_prefixes=(note_prefix,)), None
    except ImportFileError as exc:
        return None, Response({'error': str(exc)}, status=400)
    except Exception as exc:
        return None, Response({'error': f'Invalid file: {exc}'}, status=400)


def _response(created, rows_out, errors, dry_run, statuses):
    summary = {status: sum(1 for r in rows_out if r['status'] == status) for status in statuses}
    return Response({'created': created, 'errors': errors, 'dry_run': dry_run, 'summary': summary, 'rows': rows_out})


# ── members ──────────────────────────────────────────────────────────────────

def run_member_import(domain, organization, file, dry_run, by_staff=True, user=None):
    cfg = DOMAINS[domain]
    member_model, person_model, fk = cfg['member_model'], cfg['person_model'], cfg['person_fk']
    positions = {value.lower(): value for value in member_model.Position.values}
    sample = [f'01-033-02-019-000000', 'Member', '2026-01-15']
    rows, error_response = _read_rows(file, MEMBER_HEADERS, sample, MEMBER_NOTE_PREFIX)
    if error_response is not None:
        return error_response

    rows_out, errors, created = [], [], 0
    seen = set()
    import_batch = None   # opened with the first member actually added; never on a preview

    def add(row_num, status, label, detail):
        rows_out.append({'row': row_num, 'status': status, 'label': label, 'detail': detail})
        if status == 'invalid':
            errors.append({'row': row_num, 'error': detail})

    with transaction.atomic():
        for row_num, row in rows:
            rsbsa = cell_text(row[0]).strip() if row else ''
            label = rsbsa or f'Row {row_num}'
            if not rsbsa:
                add(row_num, 'invalid', label, 'RSBSA Number is required.')
                continue
            person = by_id_key(person_model.objects.all(), 'rsbsa_number', rsbsa).first()
            if person is None:
                add(row_num, 'not_found', label,
                    f'No registered {fk} has this RSBSA number. Register them first.')
                continue
            label = f'{person.full_name} ({person.rsbsa_number})'

            position_text = cell_text(row[1]).strip() if len(row) > 1 else ''
            position = positions.get(position_text.lower()) if position_text else member_model.Position.MEMBER
            if position is None:
                add(row_num, 'invalid', label,
                    f'Unknown position "{position_text}". Use one of: {", ".join(member_model.Position.values)}.')
                continue
            try:
                joined = parse_import_date(row[2] if len(row) > 2 else None)
            except ValueError as exc:
                add(row_num, 'invalid', label, str(exc))
                continue
            if joined is None:
                add(row_num, 'invalid', label, 'Date Joined is required (YYYY-MM-DD).')
                continue

            key = person.pk
            if key in seen:
                add(row_num, 'duplicate_in_file', label, 'This person appears earlier in the file; only the first row is used.')
                continue
            seen.add(key)
            if member_model.objects.filter(organization=organization, **{fk: person}).exists():
                add(row_num, 'already_member', label, 'Already a member of this organization.')
                continue

            try:
                with transaction.atomic():
                    serializer = cfg['member_serializer'](data={fk: person.pk, 'position': position, 'date_joined': joined.isoformat()})
                    if not serializer.is_valid():
                        first = next(iter(serializer.errors.values()))
                        add(row_num, 'invalid', label, str(first[0] if isinstance(first, list) else first))
                        continue
                    member = serializer.save(organization=organization)
                    _resolve_and_assign_main(
                        member,
                        member_model.objects.filter(**{f'{fk}_id': person.pk, 'is_active': True}).exclude(pk=member.pk),
                        person.barangay,
                        None,
                    )
                    if not dry_run:
                        if import_batch is None:
                            import_batch = start_batch(f'{domain}_org_members', user, label=organization.name)
                        record_created(import_batch, member)   # after the main-organization decision is saved
            except MainChoiceRequired:
                # Only staff can choose a main organization (the Add Member form ignores the choice from an
                # organization account), so an organization is told to ask staff instead.
                advice = ('add them with the Add Member form.' if by_staff
                          else "ask OMAg staff to set this person's main organization first.")
                add(row_num, 'needs_main_choice', label,
                    f'This person is in other organizations and needs a main-organization choice: {advice}')
                continue
            created += 1
            add(row_num, 'ready', label, 'Will be added.' if dry_run else 'Added.')

        if dry_run:
            transaction.set_rollback(True)
            created = 0

    return _response(created, rows_out, errors, dry_run, MEMBER_STATUSES)


# ── equipment ────────────────────────────────────────────────────────────────

EQUIPMENT_HEADERS = ['Equipment', 'Quantity', 'Condition', 'Acquisition Date (YYYY-MM-DD)', 'Remarks']
EQUIPMENT_WIDTHS = [34, 12, 20, 28, 34]
EQUIPMENT_NOTE_PREFIX = 'Valid equipment:'
EQUIPMENT_STATUSES = ('ready', 'already_exists', 'duplicate_in_file', 'invalid')


def _norm(text):
    return ' '.join(str(text or '').lower().split())


def _equipment_sample(domain):
    return [DOMAINS[domain]['equipment_model'].EquipmentName.values[0], 1, 'Good', '2026-01-15', 'Optional notes']


def run_equipment_import(domain, organization, file, dry_run, user=None):
    cfg = DOMAINS[domain]
    model = cfg['equipment_model']
    names = {_norm(v): v for v in model.EquipmentName.values}
    conditions = {_norm(v): v for v in model.Condition.values}
    rows, error_response = _read_rows(file, EQUIPMENT_HEADERS, _equipment_sample(domain), EQUIPMENT_NOTE_PREFIX)
    if error_response is not None:
        return error_response

    rows_out, errors, created = [], [], 0
    import_batch = None   # opened with the first line actually added; never on a preview

    def add(row_num, status, label, detail):
        rows_out.append({'row': row_num, 'status': status, 'label': label, 'detail': detail})
        if status == 'invalid':
            errors.append({'row': row_num, 'error': detail})

    with transaction.atomic():
        stored = {
            (e.equipment_name, _norm(e.equipment_name_other), e.condition, e.quantity, e.acquisition_date)
            for e in model.objects.filter(organization=organization)
        }
        seen = set()
        for row_num, row in rows:
            def cell(i, row=row):
                return cell_text(row[i]) if i < len(row) else ''

            typed = cell(0).strip()
            label = typed or f'Row {row_num}'
            if not typed:
                add(row_num, 'invalid', label, 'Equipment is required.')
                continue
            matched = names.get(_norm(typed))
            other_text, note = '', ''
            if matched == 'Other':
                add(row_num, 'invalid', label, 'Write what the equipment is instead of "Other".')
                continue
            if matched is None:
                if len(typed) > 100:
                    add(row_num, 'invalid', label, 'Equipment name is too long (100 characters at most).')
                    continue
                matched, other_text, note = 'Other', typed, f'Recorded as Other: {typed}. '

            quantity_text = cell(1).strip()
            if not is_plain_int(quantity_text) or int(quantity_text) < 1:
                add(row_num, 'invalid', label, 'Quantity must be a whole number of 1 or more.')
                continue
            quantity = int(quantity_text)
            condition = conditions.get(_norm(cell(2)))
            if condition is None:
                add(row_num, 'invalid', label, f'Condition must be one of: {", ".join(model.Condition.values)}.')
                continue
            try:
                acquired = parse_import_date(row[3] if len(row) > 3 else None)
            except ValueError as exc:
                add(row_num, 'invalid', label, str(exc))
                continue

            key = (matched, _norm(other_text), condition, quantity, acquired)
            if key in stored:
                add(row_num, 'already_exists', label,
                    f'{note}The organization already has this equipment with the same quantity, condition and date.')
                continue
            if key in seen:
                add(row_num, 'duplicate_in_file', label, f'{note}This row repeats an earlier row of the file.')
                continue

            with transaction.atomic():
                serializer = cfg['equipment_serializer'](data={
                    'equipment_name': matched, 'equipment_name_other': other_text, 'quantity': quantity,
                    'condition': condition, 'acquisition_date': acquired.isoformat() if acquired else None,
                    'remarks': cell(4),
                })
                if not serializer.is_valid():
                    first = next(iter(serializer.errors.values()))
                    add(row_num, 'invalid', label, str(first[0] if isinstance(first, list) else first))
                    continue
                line = serializer.save(organization=organization)
                if not dry_run:
                    if import_batch is None:
                        import_batch = start_batch(f'{domain}_org_equipment', user, label=organization.name)
                    record_created(import_batch, line)
            seen.add(key)
            created += 1
            add(row_num, 'ready', label, f'{note}{"Will be added." if dry_run else "Added."}')

        if dry_run:
            transaction.set_rollback(True)
            created = 0

    return _response(created, rows_out, errors, dry_run, EQUIPMENT_STATUSES)


class _OrgImportView(APIView):
    domain = None
    permission_classes = [IsAuthenticated, IsOwnOrgOrStaff]
    org_linked_record_type = None

    def organization(self, pk):
        return get_object_or_404(DOMAINS[self.domain]['org_model'], pk=pk)


class MemberImportTemplateView(_OrgImportView):
    def get(self, request, pk):
        self.organization(pk)
        positions = ', '.join(DOMAINS[self.domain]['member_model'].Position.values)
        return _template_response(
            f'{self.domain}_organization_members_import_template.xlsx', MEMBER_HEADERS, MEMBER_WIDTHS,
            ['01-033-02-019-000000', 'Member', '2026-01-15'], f'{MEMBER_NOTE_PREFIX} {positions}')


class MemberImportView(_OrgImportView):
    parser_classes = [MultiPartParser]

    def post(self, request, pk):
        organization = self.organization(pk)
        file = request.FILES.get('file')
        if not file:
            return Response({'error': 'No file uploaded.'}, status=400)
        dry_run = str(request.data.get('dry_run', '')).strip().lower() in ('1', 'true', 'yes')
        by_staff = request.user.role in ('ADMIN', 'EVALUATOR')
        return run_import_atomically(lambda: run_member_import(self.domain, organization, file, dry_run, by_staff, request.user))


def _for_domain(view_class, domain):
    return type(f'{domain.title()}{view_class.__name__}', (view_class,), {
        'domain': domain, 'org_linked_record_type': DOMAINS[domain]['link_type'],
    })


def member_import_view(domain):
    return _for_domain(MemberImportView, domain)


def member_import_template_view(domain):
    return _for_domain(MemberImportTemplateView, domain)


class EquipmentImportTemplateView(_OrgImportView):
    def get(self, request, pk):
        self.organization(pk)
        model = DOMAINS[self.domain]['equipment_model']
        note = (f'{EQUIPMENT_NOTE_PREFIX} {", ".join(v for v in model.EquipmentName.values if v != "Other")}. '
                f'Anything else is recorded as Other with your text. Valid conditions: {", ".join(model.Condition.values)}.')
        return _template_response(
            f'{self.domain}_organization_equipment_import_template.xlsx', EQUIPMENT_HEADERS, EQUIPMENT_WIDTHS,
            _equipment_sample(self.domain), note)


class EquipmentImportView(_OrgImportView):
    parser_classes = [MultiPartParser]

    def post(self, request, pk):
        organization = self.organization(pk)
        file = request.FILES.get('file')
        if not file:
            return Response({'error': 'No file uploaded.'}, status=400)
        dry_run = str(request.data.get('dry_run', '')).strip().lower() in ('1', 'true', 'yes')
        return run_import_atomically(lambda: run_equipment_import(self.domain, organization, file, dry_run, request.user))


def equipment_import_view(domain):
    return _for_domain(EquipmentImportView, domain)


def equipment_import_template_view(domain):
    return _for_domain(EquipmentImportTemplateView, domain)
