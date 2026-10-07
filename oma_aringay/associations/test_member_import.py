"""Importing members (already-registered people, by RSBSA number) into one organization from a file."""
import io
import itertools
from datetime import date, timedelta
from unittest import mock

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError
from openpyxl import Workbook, load_workbook
from rest_framework.test import APITestCase

from accounts.models import UserLinkedRecord
from associations.models import (
    FarmersOrganization, FarmersOrganizationMember, FisherfolkOrganization, FisherfolkOrganizationMember,
)
from farmers.models import Farmer
from fisheries.models import Fisherfolk

HEADERS = ['RSBSA Number', 'Position', 'Date Joined (YYYY-MM-DD)']
_N = itertools.count(1)
DOMAINS = {
    'farmers': dict(link='FARMERS_ORG', org=FarmersOrganization, member=FarmersOrganizationMember, fk='farmer'),
    'fisherfolk': dict(link='FISHERFOLK_ORG', org=FisherfolkOrganization, member=FisherfolkOrganizationMember, fk='fisherfolk'),
}
TODAY = date.today().isoformat()


def make_org(domain, barangay='Poblacion'):
    n = next(_N)
    return DOMAINS[domain]['org'].objects.create(
        name=f'Import Org {n}', barangay=barangay, registration_number=f'IMPM-{n:05d}', date_organized=date(2020, 1, 1),
        date_registered=date(2020, 2, 1), registration_agency='SEC', contact_person='Pat Contact',
        contact_number='09171234567')


def make_person(domain, barangay='Alaska'):
    n = next(_N)
    rsbsa = '01-033-02-019-95%04d' % n
    if domain == 'farmers':
        return Farmer.objects.create(
            rsbsa_number=rsbsa, first_name=f'Pat{n}', last_name='Importer', gender='Male', birthdate=date(1985, 1, 1),
            barangay=barangay, civil_status='Single')
    return Fisherfolk.objects.create(
        rsbsa_number=rsbsa, first_name=f'Pat{n}', last_name='Importer', gender='Male', birthdate=date(1990, 1, 1),
        barangay=barangay, fisherfolk_type='Inland')


def add_member(domain, org, person, **over):
    cfg = DOMAINS[domain]
    return cfg['member'].objects.create(organization=org, **{cfg['fk']: person}, date_joined=date(2026, 1, 1), **over)


def sheet(rows, headers=HEADERS, filename='members.xlsx'):
    book = Workbook()
    ws = book.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    stream = io.BytesIO()
    book.save(stream)
    return SimpleUploadedFile(filename, stream.getvalue(),
                              content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def make_citizen(username, link, record_id):
    user = get_user_model().objects.create_user(username=username, password='pw', role='CITIZEN', email=f'{username}@x.test')
    UserLinkedRecord.objects.create(user=user, record_type=link, record_id=record_id, is_active=True)
    return user


class MemberImportTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        cls.staff = get_user_model().objects.create_user(
            username='mi-staff', password='pw', role='ADMIN', email='mi-staff@x.test')

    def setUp(self):
        self.client.force_authenticate(self.staff)

    def url(self, domain, org, suffix='members/import/'):
        return f'/api/organizations/{domain}/{org.id}/{suffix}'

    def run_import(self, domain, org, rows, dry_run=False, headers=HEADERS):
        return self.client.post(self.url(domain, org), {'file': sheet(rows, headers), 'dry_run': 'true' if dry_run else 'false'},
                                format='multipart')

    def statuses(self, response):
        self.assertEqual(response.status_code, 200, response.content)
        return [row['status'] for row in response.data['rows']]

    def test_template_has_the_headers_and_a_note_line(self):
        for domain in DOMAINS:
            response = self.client.get(self.url(domain, make_org(domain), 'members/import/template/'))
            self.assertEqual(response.status_code, 200)
            book = load_workbook(io.BytesIO(response.content))
            ws = book.active
            self.assertEqual([c.value for c in ws[1]][:3], HEADERS)
            self.assertTrue(any(str(c.value or '').startswith('Valid positions:') for row in ws.iter_rows() for c in row))

    def test_ready_rows_create_memberships_and_first_membership_is_main(self):
        for domain, cfg in DOMAINS.items():
            org = make_org(domain)
            a, b = make_person(domain), make_person(domain)
            response = self.run_import(domain, org, [[a.rsbsa_number, 'President', TODAY], [b.rsbsa_number, '', TODAY]])
            self.assertEqual(self.statuses(response), ['ready', 'ready'], domain)
            self.assertEqual(response.data['created'], 2)
            self.assertEqual(response.data['summary']['ready'], 2)
            first = cfg['member'].objects.get(organization=org, **{cfg['fk']: a})
            second = cfg['member'].objects.get(organization=org, **{cfg['fk']: b})
            self.assertEqual((first.position, second.position), ('President', 'Member'))
            self.assertTrue(first.is_main and second.is_main)

    def test_position_and_date_problems_are_invalid(self):
        for domain, cfg in DOMAINS.items():
            org = make_org(domain)
            people = [make_person(domain) for _ in range(5)]
            future = (date.today() + timedelta(days=3)).isoformat()
            rows = [
                [people[0].rsbsa_number, 'treasurer', TODAY],      # case-insensitive position: ready
                [people[1].rsbsa_number, 'Warlord', TODAY],        # unknown position
                [people[2].rsbsa_number, 'Member', ''],            # missing date
                [people[3].rsbsa_number, 'Member', 'not a date'],  # unparseable date
                [people[4].rsbsa_number, 'Member', future],        # future date
            ]
            response = self.run_import(domain, org, rows)
            self.assertEqual(self.statuses(response), ['ready', 'invalid', 'invalid', 'invalid', 'invalid'], domain)
            self.assertIn('Treasurer', response.data['rows'][1]['detail'])
            self.assertEqual(cfg['member'].objects.filter(organization=org).get().position, 'Treasurer')

    def test_unknown_rsbsa_is_not_found(self):
        for domain in DOMAINS:
            response = self.run_import(domain, make_org(domain), [['01-033-02-019-000000', 'Member', TODAY]])
            self.assertEqual(self.statuses(response), ['not_found'])
            self.assertIn('Register them first', response.data['rows'][0]['detail'])
            self.assertEqual(response.data['created'], 0)

    def test_rsbsa_cell_is_trimmed_and_case_insensitive(self):
        for domain in DOMAINS:
            org = make_org(domain)
            person = make_person(domain)
            response = self.run_import(domain, org, [[f'  {person.rsbsa_number.lower()}  ', 'Member', TODAY]])
            self.assertEqual(self.statuses(response), ['ready'], domain)

    def test_same_person_twice_first_row_wins(self):
        for domain, cfg in DOMAINS.items():
            org = make_org(domain)
            person = make_person(domain)
            response = self.run_import(domain, org, [[person.rsbsa_number, 'Secretary', TODAY], [person.rsbsa_number, 'Auditor', TODAY]])
            self.assertEqual(self.statuses(response), ['ready', 'duplicate_in_file'])
            self.assertEqual(cfg['member'].objects.get(organization=org).position, 'Secretary')

    def test_existing_members_active_or_not_are_already_member_and_stay_as_they_are(self):
        for domain, cfg in DOMAINS.items():
            org = make_org(domain)
            active, inactive = make_person(domain), make_person(domain)
            add_member(domain, org, active)
            gone = add_member(domain, org, inactive, is_active=False)
            response = self.run_import(domain, org, [[active.rsbsa_number, 'Member', TODAY], [inactive.rsbsa_number, 'Member', TODAY]])
            self.assertEqual(self.statuses(response), ['already_member', 'already_member'], domain)
            gone.refresh_from_db()
            self.assertFalse(gone.is_active)
            self.assertEqual(cfg['member'].objects.filter(organization=org).count(), 2)

    def test_a_row_needing_the_main_choice_is_skipped_and_leaves_no_trace(self):
        for domain, cfg in DOMAINS.items():
            org = make_org(domain, barangay='Poblacion')
            other_org = make_org(domain, barangay='Basca')
            needs_choice = make_person(domain, barangay='Alaska')
            elsewhere = add_member(domain, other_org, needs_choice, is_main=False)
            fine = make_person(domain)
            response = self.run_import(domain, org, [[needs_choice.rsbsa_number, 'Member', TODAY], [fine.rsbsa_number, 'Member', TODAY]])
            self.assertEqual(self.statuses(response), ['needs_main_choice', 'ready'], domain)
            self.assertIn('Add Member', response.data['rows'][0]['detail'])
            self.assertFalse(cfg['member'].objects.filter(organization=org, **{cfg['fk']: needs_choice}).exists())
            elsewhere.refresh_from_db()
            self.assertFalse(elsewhere.is_main)
            self.assertEqual(response.data['summary']['needs_main_choice'], 1)

    def test_the_main_choice_advice_matches_who_is_importing(self):
        """An organization account cannot use the Add Member form to set a main organization, so it is told to ask staff."""
        for domain, cfg in DOMAINS.items():
            org = make_org(domain, barangay='Poblacion')
            other_org = make_org(domain, barangay='Basca')
            person = make_person(domain, barangay='Alaska')
            add_member(domain, other_org, person, is_main=False)
            rows = [[person.rsbsa_number, 'Member', TODAY]]
            staff_detail = self.run_import(domain, org, rows).data['rows'][0]['detail']
            self.assertIn('Add Member form', staff_detail)
            self.client.force_authenticate(make_citizen(f'mi-advice-{domain}', cfg['link'], org.id))
            org_detail = self.run_import(domain, org, rows).data['rows'][0]['detail']
            self.assertIn('ask OMAg staff', org_detail)
            self.assertNotIn('Add Member form', org_detail)
            self.client.force_authenticate(self.staff)

    def test_dry_run_changes_nothing_and_predicts_the_confirm(self):
        for domain, cfg in DOMAINS.items():
            org = make_org(domain, barangay='Poblacion')
            other_org = make_org(domain, barangay='Basca')
            ready, known, needs_choice = make_person(domain), make_person(domain, barangay='Alaska'), make_person(domain, barangay='Alaska')
            add_member(domain, org, known)
            add_member(domain, other_org, needs_choice, is_main=False)
            rows = [[ready.rsbsa_number, 'Member', TODAY], [known.rsbsa_number, 'Member', TODAY],
                    [needs_choice.rsbsa_number, 'Member', TODAY], ['01-033-02-019-000001', 'Member', TODAY],
                    [ready.rsbsa_number, 'Member', TODAY]]
            before = cfg['member'].objects.count()
            preview = self.run_import(domain, org, rows, dry_run=True)
            self.assertEqual(cfg['member'].objects.count(), before)
            self.assertTrue(preview.data['dry_run'])
            confirm = self.run_import(domain, org, rows)
            self.assertEqual(self.statuses(preview), self.statuses(confirm), domain)
            self.assertEqual(self.statuses(confirm), ['ready', 'already_member', 'needs_main_choice', 'not_found', 'duplicate_in_file'])
            self.assertEqual(cfg['member'].objects.count(), before + 1)

    def test_a_person_added_after_the_preview_is_already_member_in_the_confirm(self):
        for domain, cfg in DOMAINS.items():
            org = make_org(domain)
            person = make_person(domain)
            rows = [[person.rsbsa_number, 'Member', TODAY]]
            self.assertEqual(self.statuses(self.run_import(domain, org, rows, dry_run=True)), ['ready'])
            add_member(domain, org, person)
            self.assertEqual(self.statuses(self.run_import(domain, org, rows)), ['already_member'])
            self.assertEqual(cfg['member'].objects.filter(organization=org).count(), 1)

    def test_wrong_header_and_too_many_rows_are_refused(self):
        for domain in DOMAINS:
            org = make_org(domain)
            bad = self.run_import(domain, org, [['x']], headers=['Name', 'Barangay'])
            self.assertEqual(bad.status_code, 400)
            self.assertIn('error', bad.data)
            rows = [[f'01-033-02-019-9{i:05d}', 'Member', TODAY] for i in range(1001)]
            big = self.run_import(domain, org, rows)
            self.assertEqual(big.status_code, 400)

    def test_an_unexpected_database_error_saves_nothing(self):
        for domain, cfg in DOMAINS.items():
            from associations import org_imports
            serializer = org_imports.DOMAINS[domain]['member_serializer']
            real_save = serializer.save
            calls = {'n': 0}

            def second_fails(instance, *args, **kwargs):
                calls['n'] += 1
                if calls['n'] == 2:
                    raise IntegrityError('simulated')
                return real_save(instance, *args, **kwargs)

            org = make_org(domain)
            rows = [[make_person(domain).rsbsa_number, 'Member', TODAY] for _ in range(2)]
            with mock.patch.object(serializer, 'save', second_fails):
                response = self.run_import(domain, org, rows)
            self.assertEqual(response.status_code, 409, domain)
            self.assertEqual(cfg['member'].objects.filter(organization=org).count(), 0)

    def test_permissions(self):
        for domain, cfg in DOMAINS.items():
            org, other = make_org(domain), make_org(domain)
            person = make_person(domain)
            rows = [[person.rsbsa_number, 'Member', TODAY]]
            own = make_citizen(f'mi-own-{domain}', cfg['link'], org.id)
            self.client.force_authenticate(own)
            self.assertEqual(self.run_import(domain, org, rows).status_code, 200, domain)
            self.assertIn(self.run_import(domain, other, rows).status_code, (403, 404), domain)
            other_domain = 'fisherfolk' if domain == 'farmers' else 'farmers'
            foreign = make_org(other_domain)
            self.assertIn(self.run_import(other_domain, foreign, rows).status_code, (403, 404), domain)
            citizen_type = 'FARMER' if domain == 'farmers' else 'FISHERFOLK'
            person_user = make_citizen(f'mi-person-{domain}', citizen_type, person.id)
            self.client.force_authenticate(person_user)
            self.assertIn(self.run_import(domain, org, rows).status_code, (403, 404), domain)
            self.client.force_authenticate(None)
            self.assertEqual(self.run_import(domain, org, rows).status_code, 401, domain)
            self.client.force_authenticate(self.staff)

    def test_unknown_organization_is_404(self):
        response = self.client.post('/api/organizations/farmers/999999/members/import/', {'file': sheet([])}, format='multipart')
        self.assertEqual(response.status_code, 404)
