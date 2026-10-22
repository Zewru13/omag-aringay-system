"""The Excel template for the RSBSA registry import (Farmers and Fisherfolk pages).

The registry import is an UPDATE of records that already exist, matched by `rsbsa_no`; it reads the 23-column RSBSA/BFAR export
(farmers.rsbsa_import.EXPECTED_HEADERS_V2) and refuses any other header row. The template's first sheet is exactly that header row and
nothing else (a sample row would be read as data); the second sheet explains each column."""
from io import BytesIO

from django.http import HttpResponse
from rest_framework.permissions import IsAuthenticated
from rest_framework.views import APIView

from accounts.permissions import IsAdminOrEvaluator
from core.exports import SafeWorkbook

from .rsbsa_import import EXPECTED_HEADERS_V2

COLUMN_HELP = {
    'rsbsa_no': ('The RSBSA number', 'Matches the existing farmer or fisherfolk. A row with no match is reported, never created.'),
    'control_no': ('The control number of the export', 'Any text.'),
    'first_name': ('First name', 'Text.'),
    'middle_name': ('Middle name', 'Text, or NULL / blank.'),
    'last_name': ('Last name', 'Text.'),
    'ext_name': ('Name extension (Jr., Sr., III)', 'Text, or NULL / blank.'),
    'farmer_address_bgy': ('Barangay', 'An Aringay barangay, for example Alaska or Pangao-aoan East.'),
    'farmer_address_mun': ('Municipality', 'Aringay.'),
    'farmer_address_prv': ('Province', 'La Union.'),
    'birthday': ('Date of birth', 'YYYY-MM-DD (DD/MM/YYYY is also read).'),
    'gender': ('Gender', 'Male or Female.'),
    'contact_num': ('Mobile number', '09XXXXXXXXX (a 10-digit number starting with 9 is read as 09...). Blank is fine.'),
    'cropname': ('Crop or commodity', 'Text. Read but not used to change records.'),
    'crop_area': ('Area of the crop in hectares', 'A number. Read but not used to change records.'),
    'verified_area(for_testing)': ('Verified area', 'A number. Read but not used to change records.'),
    'farmer': ('Is a farmer', 'YES or NO.'),
    'farmworker': ('Is a farmworker', 'YES or NO.'),
    'fisherfolk': ('Is a fisherfolk', 'YES or NO. At least one of farmer, farmworker, fisherfolk must be YES.'),
    'agency': ('Agency that registered the person', 'Text.'),
    'farm_type': ('Farm type', 'Text.'),
    'date_created': ('Date the record was created', 'YYYY-MM-DD.'),
    'kyc_updated': ('KYC update date', 'Text or a date.'),
    'parcel_updated': ('Parcel update date', 'Text or a date.'),
}


class RsbsaImportTemplateView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def get(self, request):
        from openpyxl.styles import Font, PatternFill

        book = SafeWorkbook()
        data = book.active
        data.title = 'RSBSA Registry'
        fill = PatternFill(start_color='1E4620', end_color='1E4620', fill_type='solid')
        font = Font(color='FFFFFF', bold=True)
        for column, header in enumerate(EXPECTED_HEADERS_V2, 1):
            cell = data.cell(row=1, column=column, value=header)
            cell.fill, cell.font = fill, font
            data.column_dimensions[cell.column_letter].width = max(14, len(header) + 2)

        help_sheet = book.create_sheet('How to fill this in')
        help_sheet.append(['This import updates farmers and fisherfolk who are already registered, matched by rsbsa_no. '
                           'Paste the rows of the RSBSA/BFAR export under the header row of the first sheet. '
                           'Do not rename, reorder or delete the columns: the import refuses a file whose header row differs.'])
        help_sheet.append([])
        help_sheet.append(['Column', 'What it is', 'What to enter'])
        for cell in help_sheet[3]:
            cell.fill, cell.font = fill, font
        for header in EXPECTED_HEADERS_V2:
            what, how = COLUMN_HELP.get(header, ('', ''))
            help_sheet.append([header, what, how])
        for letter, width in (('A', 28), ('B', 38), ('C', 80)):
            help_sheet.column_dimensions[letter].width = width

        buffer = BytesIO()
        book.save(buffer)
        response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="rsbsa_registry_import_template.xlsx"'
        return response
