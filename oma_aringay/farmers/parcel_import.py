"""File imports for farm parcels.

Two files share one set of crop columns:

* the crops file read by the Add Farm Parcel wizard (`ParcelCropsParseView`): it only checks the file and returns the crops,
  which the wizard adds to its own list. Nothing is saved, so there is nothing to undo;
* the parcels file imported from a farmer's profile (`FarmParcelImportView`): one row per CROP, the parcel columns repeated.
  Rows that share a parcel number are one parcel with several crops; a parcel without crops is one row with the crop cells
  blank; rows with no parcel number are separate parcels. Preview and commit run the same checks, a commit is one
  transaction, and what it creates is remembered as an undoable batch (each parcel, then its crops).
"""
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from rest_framework.parsers import MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from accounts.permissions import IsAdminOrEvaluator
from core.constants import BARANGAY_CHOICES
from core.exports import SafeWorkbook
from core.import_batches import record_created, start_batch
from core.utils import ImportFileError, cell_text, read_template_rows, run_import_atomically

from .models import FarmCrop, FarmParcel, Farmer, ParcelClassification
from .rsbsa_import import normalize_barangay
from .serializers import FarmCropSerializer, FarmParcelSerializer

PARCEL_COLUMNS = [
    'Parcel Number', 'Area (ha)', 'Barangay', 'Land Ownership', 'Land Owner Name', 'Land Owner RSBSA Number',
    'Tax Dec Number', 'Farm Type', 'Classifications', 'Cropping Schedule',
]
CROP_HEADERS = [
    'Crop Name', 'Crop Area (ha)', 'Main Crop (Yes/No)', 'Variety Used', 'Seed Source', 'Qty of Seeds', 'Seeds Unit', 'No. of Trees',
]
PARCEL_HEADERS = PARCEL_COLUMNS + CROP_HEADERS
PARCEL_WIDTHS = [26, 12, 20, 16, 24, 26, 26, 20, 34, 20, 24, 14, 14, 20, 18, 12, 12, 12]
CROP_WIDTHS = PARCEL_WIDTHS[len(PARCEL_COLUMNS):]

PARCEL_SAMPLE = ['01-33-02-005-001-02-003', '2.50', 'Alaska', 'Owner', '', '', '', 'Irrigated', 'Organic Farm', 'Double Cropping']
CROP_SAMPLE = ['Palay (Rice)', '2.5', 'Yes', 'NSIC Rc 222', 'Certified Seed', '40', 'kg', '']
PARCEL_NOTE_PREFIX = 'How to fill this in:'
CROP_NOTE_PREFIX = 'How to fill this in:'
MAX_PARCEL_ROWS = 1000
MAX_CROP_ROWS = 100

SEED_SOURCES = [value for value, _ in FarmCrop.SEED_SOURCE_CHOICES]
SEED_UNITS = [value for value, _ in FarmCrop.SEEDS_UNIT_CHOICES]
CROP_NAMES = [value for value, _ in FarmCrop.CropName.choices if value != 'Other']
OWNERSHIPS = [value for value, _ in FarmParcel.LandOwnership.choices]
FARM_TYPES = [value for value, _ in FarmParcel.FARM_TYPE_CHOICES]
CROPPING_SCHEDULES = [value for value, _ in FarmParcel.CROPPING_SCHEDULE_CHOICES]
BARANGAY_NAMES = [name for name, _ in BARANGAY_CHOICES]

SUMMARY_STATUSES = ('ready', 'duplicate_in_file', 'duplicate_in_database', 'invalid')


# ── reading cells ──────────────────────────────────────────────────────────────────────────────────────────────────

def _text(row, index):
    return cell_text(row[index]).strip() if index < len(row) else ''


def _yes_no(text, label):
    lowered = text.lower()
    if lowered in ('', 'no', 'n', 'false', '0'):
        return False
    if lowered in ('yes', 'y', 'true', '1'):
        return True
    raise ValueError(f'{label} must be Yes or No (found "{text}").')


def _choice(text, choices, label):
    """The canonical spelling of `text` among `choices` ('' stays ''); anything else is an error listing the choices."""
    if not text:
        return ''
    for choice in choices:
        if choice.lower() == text.lower():
            return choice
    raise ValueError(f'{label} "{text}" is not valid. Use one of: {", ".join(choices)}.')


def _decimal(text, label, *, places, minimum_exclusive=None):
    try:
        value = Decimal(text)
    except InvalidOperation:
        raise ValueError(f'{label} "{text}" is not a number.')
    if not value.is_finite():
        raise ValueError(f'{label} "{text}" is not a number.')
    if minimum_exclusive is not None and value <= minimum_exclusive:
        raise ValueError(f'{label} must be more than {minimum_exclusive}.')
    if value < 0:
        raise ValueError(f'{label} cannot be negative.')
    if -value.as_tuple().exponent > places:
        raise ValueError(f'{label} can have at most {places} decimal places.')
    return value


def _serializer_errors(serializer):
    return '; '.join(f'{field}: {"; ".join(str(m) for m in messages)}' for field, messages in serializer.errors.items())


def parse_crop_cells(row, offset=0):
    """One crop from a row of cells (`offset` is where the crop columns start). Blank crop cells mean "no crop" and give
    None; a crop with a missing name or area, or a value that is not valid, raises ValueError with a sentence."""
    cells = [_text(row, offset + i) for i in range(len(CROP_HEADERS))]
    name, area, main, variety, seed, qty, unit, trees = cells
    if not any(cells):
        return None
    if not name:
        raise ValueError('Crop Name is required when any other crop cell is filled in.')
    if not area:
        raise ValueError('Crop Area (ha) is required.')
    crop_name = _choice(name, CROP_NAMES, 'Crop Name') if name.lower() in {c.lower() for c in CROP_NAMES} else 'Other'
    area_value = _decimal(area, 'Crop area', places=4, minimum_exclusive=Decimal(0))
    seed_source = _choice(seed, SEED_SOURCES, 'Seed source')
    quantity = _decimal(qty, 'Qty of seeds', places=2) if qty else None
    seeds_unit = _choice(unit, SEED_UNITS, 'Seeds unit') or ('kg' if quantity is not None else '')
    number_of_trees = None
    if trees:
        try:
            number_of_trees = int(Decimal(trees))
        except (InvalidOperation, ValueError):
            raise ValueError(f'No. of trees "{trees}" is not a whole number.')
        if number_of_trees < 0:
            raise ValueError('No. of trees cannot be negative.')
    crop = {
        'crop_name': crop_name,
        'crop_name_other': name if crop_name == 'Other' else '',
        'area_hectares': area_value,
        'is_main_crop': _yes_no(main, 'Main Crop'),
        'variety_used': variety,
        'seed_source': seed_source,
        'quantity_of_seeds_used': quantity,
        'seeds_unit': seeds_unit,
        'number_of_trees': number_of_trees,
    }
    check = FarmCropSerializer(data={**crop, 'area_hectares': str(area_value),
                                     'quantity_of_seeds_used': None if quantity is None else str(quantity)})
    if not check.is_valid():
        raise ValueError(_serializer_errors(check))
    return crop


def parse_parcel_cells(row):
    """The parcel columns of a row as a dict for FarmParcelSerializer, or ValueError with a sentence."""
    number, area, barangay, ownership, owner, owner_rsbsa, tax, farm_type, classifications, schedule = (
        _text(row, i) for i in range(len(PARCEL_COLUMNS)))
    if not area:
        raise ValueError('Area (ha) is required.')
    land_ownership = _choice(ownership, OWNERSHIPS, 'Land ownership') or 'Owner'
    if land_ownership in ('Tenant', 'Lessee') and not owner:
        raise ValueError('Land Owner Name is required for a Tenant or Lessee parcel.')
    if not barangay:
        raise ValueError('Barangay is required.')
    canonical_barangay = normalize_barangay(barangay)
    if canonical_barangay not in BARANGAY_NAMES:
        raise ValueError(f'Barangay "{barangay}" is not one of the Aringay barangays.')
    parcel = {
        'parcel_number': number,
        'area_hectares': str(_decimal(area, 'Area (ha)', places=2, minimum_exclusive=Decimal(0))),
        'barangay': canonical_barangay,
        'land_ownership': land_ownership,
        'land_owner_name': owner if land_ownership in ('Tenant', 'Lessee') else '',
        'land_owner_rsbsa_number': owner_rsbsa if land_ownership in ('Tenant', 'Lessee') else '',
        'tax_declaration_number': tax,
        'farm_type': _choice(farm_type, FARM_TYPES, 'Farm type'),
        'classifications': [c.strip() for c in classifications.replace(';', ',').split(',') if c.strip()],
        'cropping_schedule': _choice(schedule, CROPPING_SCHEDULES, 'Cropping schedule'),
    }
    check = FarmParcelSerializer(data=parcel)
    if not check.is_valid():
        raise ValueError(_serializer_errors(check))
    return parcel


# ── templates ──────────────────────────────────────────────────────────────────────────────────────────────────────

def _template_response(filename, headers, widths, sample_row, note):
    from openpyxl.styles import Font, PatternFill

    book = SafeWorkbook()
    sheet = book.active
    sheet.title = 'Import'
    fill = PatternFill(start_color='1E4620', end_color='1E4620', fill_type='solid')
    font = Font(color='FFFFFF', bold=True)
    for column, (header, width) in enumerate(zip(headers, widths), 1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.fill = fill
        cell.font = font
        sheet.column_dimensions[cell.column_letter].width = width
    for column, value in enumerate(sample_row, 1):
        sheet.cell(row=2, column=column, value=value)
    sheet.cell(row=4, column=1, value=note)
    buffer = BytesIO()
    book.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


CROP_NOTE = (
    f'{CROP_NOTE_PREFIX} one row per crop. Crop Name: {", ".join(CROP_NAMES)} (anything else is recorded as Other with your text). '
    f'Main Crop: Yes for the one main crop, otherwise No. Seed Source: {", ".join(SEED_SOURCES)}. Seeds Unit: {", ".join(SEED_UNITS)}. '
    'No. of Trees is for tree crops only. Delete the sample row.'
)
def parcel_note():
    """The template's instructions; the Classifications line names what Classification Settings offers right now."""
    offered = ', '.join(ParcelClassification.objects.filter(is_active=True).order_by('name').values_list('name', flat=True)) or 'none set up yet'
    return (
        f'{PARCEL_NOTE_PREFIX} one row per CROP, with the parcel columns repeated; rows with the same Parcel Number are one parcel. '
        'A parcel with no crops is one row with the crop columns left blank. Rows without a Parcel Number are separate parcels. '
        f'Land Ownership: {", ".join(OWNERSHIPS)} (a Tenant or Lessee also needs the Land Owner Name). '
        f'Farm Type: {", ".join(FARM_TYPES)}. Classifications: any of {offered}, separated by commas (leave blank for none). '
        f'Cropping Schedule: {", ".join(CROPPING_SCHEDULES)}. Barangay must be an Aringay barangay. '
        f'Crop columns: {CROP_NOTE[len(CROP_NOTE_PREFIX) + 1:]}'
    )


class ParcelCropsTemplateView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def get(self, request):
        return _template_response('farm_parcel_crops_import_template.xlsx', CROP_HEADERS, CROP_WIDTHS, CROP_SAMPLE, CROP_NOTE)


class ParcelImportTemplateView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def get(self, request):
        return _template_response('farm_parcels_import_template.xlsx', PARCEL_HEADERS, PARCEL_WIDTHS, PARCEL_SAMPLE + CROP_SAMPLE, parcel_note())


def _read(file, headers, sample, prefix, max_rows):
    try:
        return read_template_rows(file, file.name, headers, sample_row=sample, note_prefixes=(prefix,), max_rows=max_rows), None
    except ImportFileError as exc:
        return None, Response({'error': str(exc)}, status=400)
    except Exception as exc:
        return None, Response({'error': f'Invalid file: {exc}'}, status=400)


# ── the wizard's crops file: checked and handed back, nothing saved ──────────────────────────────────────────────────

class ParcelCropsParseView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    parser_classes = [MultiPartParser]

    def post(self, request):
        file = request.FILES.get('file')
        if not file:
            return Response({'error': 'No file uploaded.'}, status=400)
        rows, failure = _read(file, CROP_HEADERS, CROP_SAMPLE, CROP_NOTE_PREFIX, MAX_CROP_ROWS)
        if failure:
            return failure
        crops, errors = [], []
        for number, row in rows:
            try:
                crop = parse_crop_cells(row)
            except ValueError as exc:
                errors.append({'row': number, 'error': str(exc)})
                continue
            if crop:
                crops.append((number, crop))
        mains = [number for number, crop in crops if crop['is_main_crop']]
        if len(mains) > 1:
            errors.append({'row': mains[1], 'error': 'A parcel has only one main crop; more than one row says Yes.'})
        if errors:
            errors.sort(key=lambda e: e['row'])
            return Response({'error': ' '.join(f'Row {e["row"]}: {e["error"]}' for e in errors), 'errors': errors}, status=400)
        return Response({'crops': [{
            **crop, 'area_hectares': float(crop['area_hectares']),
            'quantity_of_seeds_used': None if crop['quantity_of_seeds_used'] is None else float(crop['quantity_of_seeds_used']),
        } for _, crop in crops]})


# ── the parcels file ────────────────────────────────────────────────────────────────────────────────────────────────

def _group_rows(rows):
    """Parse every row and gather the rows of one parcel: {key: [item]} in file order. An item is
    {'row', 'parcel', 'crop', 'error'}."""
    groups = {}
    for number, row in rows:
        item = {'row': number, 'parcel': None, 'crop': None, 'error': ''}
        try:
            item['parcel'] = parse_parcel_cells(row)
            item['crop'] = parse_crop_cells(row, offset=len(PARCEL_COLUMNS))
        except ValueError as exc:
            item['error'] = str(exc)
        number_text = _text(row, 0)
        key = number_text.lower() if number_text else f'row-{number}'
        groups.setdefault(key, []).append(item)
    return groups


def _judge(farmer, items):
    """(status, detail) for one parcel's rows."""
    rows_text = ', '.join(str(i['row']) for i in items)
    if any(i['error'] for i in items):
        return 'invalid', None
    first = items[0]['parcel']
    if any(i['parcel'] != first for i in items[1:]):
        return 'duplicate_in_file', f'Rows {rows_text} are the same parcel but give different details for it. Nothing was imported for it.'
    if sum(1 for i in items if i['crop'] and i['crop']['is_main_crop']) > 1:
        return 'invalid', f'A parcel has only one main crop; more than one of rows {rows_text} says Yes.'
    number = first['parcel_number']
    if number and FarmParcel.objects.filter(farmer=farmer, parcel_number__iexact=number).exists():
        return 'duplicate_in_database', f'This farmer already has parcel {number}. Not imported again.'
    return 'ready', None


def run_parcel_import(farmer, file, dry_run, user):
    rows, failure = _read(file, PARCEL_HEADERS, PARCEL_SAMPLE + CROP_SAMPLE, PARCEL_NOTE_PREFIX, MAX_PARCEL_ROWS)
    if failure:
        return failure

    rows_out, errors = [], []
    created = 0
    batch = None
    for items in _group_rows(rows).values():
        status, detail = _judge(farmer, items)
        if status == 'ready' and not dry_run:
            parcel_data = items[0]['parcel']
            serializer = FarmParcelSerializer(data=parcel_data)
            serializer.is_valid(raise_exception=True)
            parcel = serializer.save(farmer=farmer)
            crops = []
            for item in items:
                if item['crop']:
                    crop_data = item['crop']
                    crop_serializer = FarmCropSerializer(data={**crop_data, 'area_hectares': str(crop_data['area_hectares']),
                                                               'quantity_of_seeds_used': None if crop_data['quantity_of_seeds_used'] is None
                                                               else str(crop_data['quantity_of_seeds_used'])})
                    crop_serializer.is_valid(raise_exception=True)
                    crops.append(crop_serializer.save(farm_parcel=parcel))
            if batch is None:
                batch = start_batch('farm_parcels', user, label=file.name)
            # Parcel first: an undo then judges the parcel while its crops are still there (removing a crop changes the
            # parcel's commodity, which would make it look edited) and takes the crops with it. Recorded after the crops
            # were saved, which is when the parcel's commodity is set.
            record_created(batch, parcel)
            for crop in crops:
                record_created(batch, crop)
            created += 1
        for item in items:
            number = item['parcel']['parcel_number'] if item['parcel'] else ''
            crop = item['crop']['crop_name'] if item['crop'] else ''
            label = f'{number or "(no parcel number)"} — {crop or "no crop"}'
            if status == 'invalid':
                row_detail = item['error'] or detail or 'Another row of this parcel has a problem, so the parcel was not imported.'
                errors.append({'row': item['row'], 'error': row_detail})
            else:
                row_detail = detail or ('Will be created.' if dry_run else 'Created.')
            rows_out.append({'row': item['row'], 'status': status, 'label': label, 'detail': row_detail})
    rows_out.sort(key=lambda r: r['row'])
    errors.sort(key=lambda e: e['row'])
    summary = {status: sum(1 for r in rows_out if r['status'] == status) for status in SUMMARY_STATUSES}
    return Response({'created': created, 'errors': errors, 'dry_run': dry_run, 'summary': summary, 'rows': rows_out})


class FarmParcelImportView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    parser_classes = [MultiPartParser]

    def post(self, request, farmer_id):
        farmer = get_object_or_404(Farmer, pk=farmer_id)
        file = request.FILES.get('file')
        if not file:
            return Response({'error': 'No file uploaded.'}, status=400)
        dry_run = str(request.data.get('dry_run', '')).strip().lower() in ('1', 'true', 'yes')
        return run_import_atomically(lambda: run_parcel_import(farmer, file, dry_run, request.user))
