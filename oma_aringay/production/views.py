from collections import defaultdict
from decimal import Decimal
from io import BytesIO

from django.db.models import Avg, Case, CharField, Count, Exists, IntegerField, OuterRef, Q, Sum, Value, When
from django.shortcuts import get_object_or_404
from django.utils import timezone
from rest_framework import generics, status as drf_status
from rest_framework.parsers import MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from accounts.permissions import IsAdmin, IsAdminOrEvaluator
from core.constants import BARANGAY_CHOICES
from farmers.models import FarmParcel, Farmer
from fisheries.models import Fisherfolk

from .models import (
    BarangayCropSummary,
    CropActivityLog, CropActivityStage, CropCategory, CropDisasterReport,
    CropMonitoringRecord, CropRemark,
    CropSubcategory, CropType, DamageReport, DisasterResponseRule,
    FarmInputLog, HarvestRecord, LaborOperationLog, PestDiseaseReport, PlantingRecord,
)
from .serializers import (
    BarangayCropSummarySerializer,
    CropActivityLogSerializer,
    CropActivityStageSerializer,
    CropCategorySerializer,
    CropDisasterReportSerializer,
    CropMonitoringDetailSerializer,
    CropMonitoringRecordSerializer,
    CropRemarkSerializer,
    MonitoringRemarkSerializer,
    CropSubcategorySerializer,
    CropTypeSerializer,
    DamageReportSerializer,
    FarmInputLogSerializer,
    HarvestRecordCreateSerializer,
    HarvestRecordSerializer,
    LaborOperationLogSerializer,
    PestDiseaseReportSerializer,
    PlantingRecordCreateSerializer,
    PlantingRecordSerializer,
)


class HarvestRecordListView(generics.ListAPIView):
    serializer_class = HarvestRecordSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    pagination_class = None

    def get_queryset(self):
        queryset = HarvestRecord.objects.select_related(
            'farm_parcel',
            'farm_parcel__farmer',
            'recorded_by',
        ).order_by('-date_recorded')

        year = self.request.query_params.get('year')
        season = self.request.query_params.get('season')
        barangay = self.request.query_params.get('barangay')
        commodity = self.request.query_params.get('commodity')

        if year:
            queryset = queryset.filter(year=year)

        if season:
            queryset = queryset.filter(season__iexact=season)

        if barangay:
            queryset = queryset.filter(farm_parcel__barangay__iexact=barangay)

        if commodity:
            queryset = queryset.filter(
                Q(crop_name__icontains=commodity)
                | Q(farm_parcel__commodity_name__icontains=commodity)
                | Q(farm_parcel__commodity_type__icontains=commodity)
            )

        return queryset


class HarvestRecordCreateView(generics.CreateAPIView):
    queryset = HarvestRecord.objects.all()
    serializer_class = HarvestRecordCreateSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def perform_create(self, serializer):
        serializer.save(recorded_by=self.request.user)


class HarvestRecordDetailView(generics.RetrieveUpdateDestroyAPIView):
    queryset = HarvestRecord.objects.select_related(
        'farm_parcel',
        'farm_parcel__farmer',
        'recorded_by',
    )
    serializer_class = HarvestRecordSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    http_method_names = ['get', 'patch', 'delete']
    lookup_field = 'id'

    def destroy(self, request, *args, **kwargs):
        if request.user.role != 'ADMIN':
            return Response(
                {'detail': 'Only administrators can delete harvest records.'},
                status=drf_status.HTTP_403_FORBIDDEN,
            )
        harvest = self.get_object()
        monitoring = harvest.monitoring_record
        response = super().destroy(request, *args, **kwargs)
        if monitoring is not None:
            monitoring.status = CropMonitoringRecord.Status.ACTIVE
            monitoring.actual_yield_kg = None
            monitoring.actual_harvest_date = None
            monitoring.yield_variance_kg = None
            monitoring.yield_variance_pct = None
            monitoring.save()
        return response


class ProductionDashboardView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def get(self, request):
        current_year = timezone.now().year
        harvests_this_year = HarvestRecord.objects.filter(year=current_year)

        total_volume = harvests_this_year.aggregate(
            total=Sum('volume_kg'),
        )['total'] or 0

        top_commodity_record = harvests_this_year.values('crop_name').annotate(
            total_volume=Sum('volume_kg'),
        ).order_by('-total_volume').first()

        return Response({
            'total_farmers': Farmer.objects.count(),
            'active_farms': FarmParcel.objects.filter(farm_status='Active').count(),
            'inactive_farms': FarmParcel.objects.filter(farm_status='Inactive').count(),
            'total_volume_kg_this_year': total_volume,
            'top_commodity': top_commodity_record['crop_name'] if top_commodity_record else None,
            'total_fisherfolk': Fisherfolk.objects.count(),
        })


class AgriDashboardView(APIView):
    """Dashboard stats for the Agri module: associations, crop records, harvest calendar, distribution.
    Pass ?domain=fisheries to get fisheries stocking stats instead."""
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def get(self, request):
        domain = request.query_params.get('domain', 'crops').strip().lower()
        if domain == 'fisheries':
            return self._get_fisheries(request)
        return self._get_crops(request)

    def _get_fisheries(self, request):
        from associations.models import FisherfolkAssociation
        from fisheries.models import StockingRecord
        from datetime import date, timedelta

        today      = date.today()
        two_weeks  = today + timedelta(days=14)
        year_param = request.query_params.get('year',   '').strip()

        total_assoc = FisherfolkAssociation.objects.count()

        base_active = StockingRecord.objects.filter(
            status__in=[StockingRecord.Status.STOCKED, StockingRecord.Status.FOR_HARVEST]
        )
        if year_param:
            base_active = base_active.filter(year=year_param)

        total_stocking    = base_active.count()
        nearing_harvest   = base_active.filter(
            expected_harvest_date__lte=two_weeks,
            expected_harvest_date__gte=today,
        ).count()

        top_barangay_row = (
            base_active
            .exclude(pond_location_barangay='')
            .values('pond_location_barangay')
            .annotate(total_area=Sum('area_hectares'))
            .order_by('-total_area')
            .first()
        )
        top_barangay = top_barangay_row['pond_location_barangay'] if top_barangay_row else '—'

        # Harvest calendar from stocking expected_harvest_date
        upcoming_qs = (
            base_active
            .filter(
                expected_harvest_date__gte=today,
                expected_harvest_date__lte=today + timedelta(days=90),
            )
            .values_list('expected_harvest_date', flat=True)
            .distinct()
        )
        overdue_qs = (
            base_active
            .filter(expected_harvest_date__lt=today)
            .values_list('expected_harvest_date', flat=True)
            .distinct()
        )
        harvest_dates = (
            [{'date': str(d), 'type': 'upcoming'} for d in upcoming_qs]
            + [{'date': str(d), 'type': 'overdue'}  for d in overdue_qs]
        )

        # Distribution by species + barangay
        all_stocking = StockingRecord.objects.select_related('species')
        if year_param:
            all_stocking = all_stocking.filter(year=year_param)

        dist_qs = (
            all_stocking
            .values('pond_location_barangay', 'species__species_name')
            .annotate(count=Count('id'))
            .order_by('species__species_name')
        )
        distribution_detail = [
            {
                'barangay': row['pond_location_barangay'] or 'Unknown',
                'crop':     row['species__species_name']  or 'Unknown',
                'count':    row['count'],
            }
            for row in dist_qs
        ]

        available_years = sorted(
            StockingRecord.objects.values_list('year', flat=True).distinct(),
            reverse=True,
        )

        return Response({
            'total_associations': total_assoc,
            'total_crop_records': total_stocking,
            'nearing_harvest':    nearing_harvest,
            'top_barangay':       top_barangay,
            'harvest_dates':      harvest_dates,
            'harvest_calendar':   harvest_dates,
            'crop_distribution':  [],
            'crop_distribution_detail': distribution_detail,
            'available_years':    available_years,
        })

    def _get_crops(self, request):
        from associations.models import FarmersAssociation
        from datetime import date, timedelta

        today = date.today()
        week_later = today + timedelta(days=7)

        season_param = request.query_params.get('season', '').strip()
        year_param   = request.query_params.get('year',   '').strip()

        def apply_filters(qs):
            if season_param:
                qs = qs.filter(season__iexact=season_param)
            if year_param:
                qs = qs.filter(year=year_param)
            return qs

        total_associations = FarmersAssociation.objects.count()

        base_active = apply_filters(
            CropMonitoringRecord.objects.filter(status=CropMonitoringRecord.Status.ACTIVE)
        )
        total_crop_records = base_active.count()

        nearing_harvest = base_active.filter(
            expected_harvest_date__gte=today,
            expected_harvest_date__lte=week_later,
        ).count()

        # Top barangay by total area planted
        top_barangay_row = (
            base_active
            .values('farm_location_barangay')
            .annotate(total_area=Sum('area_planted_ha'))
            .exclude(farm_location_barangay='')
            .order_by('-total_area')
            .first()
        )
        if not top_barangay_row:
            top_barangay_row = (
                base_active
                .values('farm_parcel__barangay')
                .annotate(total_area=Sum('area_planted_ha'))
                .exclude(farm_parcel__barangay='')
                .order_by('-total_area')
                .first()
            )
            top_barangay = top_barangay_row['farm_parcel__barangay'] if top_barangay_row else '—'
        else:
            top_barangay = top_barangay_row['farm_location_barangay']

        # Harvest dates: upcoming (next 90 days) + overdue (past, still active)
        upcoming_qs = (
            base_active
            .filter(
                expected_harvest_date__gte=today,
                expected_harvest_date__lte=today + timedelta(days=90),
            )
            .values_list('expected_harvest_date', flat=True)
            .distinct()
        )
        overdue_qs = (
            base_active
            .filter(expected_harvest_date__lt=today)
            .values_list('expected_harvest_date', flat=True)
            .distinct()
        )
        harvest_dates = (
            [{'date': str(d), 'type': 'upcoming'} for d in upcoming_qs]
            + [{'date': str(d), 'type': 'overdue'}  for d in overdue_qs]
        )
        harvest_calendar = harvest_dates

        # Crop distribution (all statuses, filtered by season/year)
        dist_base = apply_filters(CropMonitoringRecord.objects.all())
        dist_qs = (
            dist_base
            .values('farm_location_barangay', 'farm_parcel__barangay')
            .annotate(count=Count('id'), area_ha=Sum('area_planted_ha'))
            .order_by('farm_location_barangay')
        )
        crop_distribution = [
            {
                'barangay': (row['farm_location_barangay'] or row['farm_parcel__barangay'] or 'Unknown'),
                'count': row['count'],
                'area_ha': float(row['area_ha'] or 0),
            }
            for row in dist_qs
        ]

        # Crop distribution detail: barangay + crop grouped
        detail_qs = (
            dist_base
            .values('farm_location_barangay', 'farm_parcel__barangay', 'farmers_association__barangay', 'crop_type__name')
            .annotate(count=Count('id'))
            .order_by('crop_type__name')
        )
        crop_distribution_detail = [
            {
                'barangay': (row['farm_location_barangay']
                             or row['farm_parcel__barangay']
                             or row['farmers_association__barangay']
                             or 'Unknown'),
                'crop': row['crop_type__name'] or 'Unknown',
                'count': row['count'],
            }
            for row in detail_qs
        ]

        # Available years for the filter dropdowns
        available_years = sorted(
            CropMonitoringRecord.objects.values_list('year', flat=True).distinct(),
            reverse=True,
        )

        return Response({
            'total_associations': total_associations,
            'total_crop_records': total_crop_records,
            'nearing_harvest': nearing_harvest,
            'top_barangay': top_barangay,
            'harvest_dates': harvest_dates,
            'harvest_calendar': harvest_calendar,
            'crop_distribution': crop_distribution,
            'crop_distribution_detail': crop_distribution_detail,
            'available_years': available_years,
        })


class BarangaySummaryView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def get(self, request):
        queryset = HarvestRecord.objects.all()
        year = request.query_params.get('year')
        season = request.query_params.get('season')

        if year:
            queryset = queryset.filter(year=year)

        if season:
            queryset = queryset.filter(season__iexact=season)

        data = queryset.values(
            'farm_parcel__barangay',
        ).annotate(
            total_volume_kg=Sum('volume_kg'),
            total_area_hectares=Sum('area_harvested_hectares'),
            number_of_farmers=Count('farm_parcel__farmer', distinct=True),
        ).order_by('-total_volume_kg')

        return Response([
            {
                'barangay_name': item['farm_parcel__barangay'],
                'total_volume_kg': item['total_volume_kg'] or 0,
                'total_area_hectares': item['total_area_hectares'] or 0,
                'number_of_farmers': item['number_of_farmers'],
            }
            for item in data
        ])


class CommodityRankingView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def get(self, request):
        queryset = HarvestRecord.objects.all()
        year = request.query_params.get('year')

        if year:
            queryset = queryset.filter(year=year)

        data = queryset.values(
            'crop_name',
        ).annotate(
            total_volume_kg=Sum('volume_kg'),
            total_area_hectares=Sum('area_harvested_hectares'),
            number_of_farms=Count('farm_parcel', distinct=True),
        ).order_by('-total_volume_kg')

        return Response([
            {
                'crop_name': item['crop_name'],
                'total_volume_kg': item['total_volume_kg'] or 0,
                'total_area_hectares': item['total_area_hectares'] or 0,
                'number_of_farms': item['number_of_farms'],
            }
            for item in data
        ])


class SeasonalComparisonView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def get(self, request):
        queryset = HarvestRecord.objects.all()
        year = request.query_params.get('year')

        if year:
            queryset = queryset.filter(year=year)

        dry_season = queryset.filter(season='Dry').aggregate(
            total_volume_kg=Sum('volume_kg'),
            farm_count=Count('farm_parcel', distinct=True),
        )
        wet_season = queryset.filter(season='Wet').aggregate(
            total_volume_kg=Sum('volume_kg'),
            farm_count=Count('farm_parcel', distinct=True),
        )

        return Response({
            'dry_season': {
                'total_volume_kg': dry_season['total_volume_kg'] or 0,
                'farm_count': dry_season['farm_count'],
            },
            'wet_season': {
                'total_volume_kg': wet_season['total_volume_kg'] or 0,
                'farm_count': wet_season['farm_count'],
            },
        })


class PlantingRecordListCreateView(generics.ListCreateAPIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    pagination_class = None

    def get_serializer_class(self):
        if self.request.method == 'POST':
            return PlantingRecordCreateSerializer
        return PlantingRecordSerializer

    def get_queryset(self):
        queryset = PlantingRecord.objects.select_related(
            'farm_parcel', 'farm_parcel__farmer',
        ).order_by('-date_recorded')
        year = self.request.query_params.get('year')
        season = self.request.query_params.get('season')
        barangay = self.request.query_params.get('barangay')
        if year:
            queryset = queryset.filter(year=year)
        if season:
            queryset = queryset.filter(season__iexact=season)
        if barangay:
            queryset = queryset.filter(farm_parcel__barangay__iexact=barangay)
        return queryset

    def perform_create(self, serializer):
        serializer.save(recorded_by=self.request.user)


class PlantingRecordDetailView(generics.RetrieveUpdateDestroyAPIView):
    queryset = PlantingRecord.objects.select_related('farm_parcel', 'farm_parcel__farmer')
    serializer_class = PlantingRecordSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    http_method_names = ['get', 'patch', 'delete']
    lookup_field = 'pk'


class DamageReportListCreateView(generics.ListCreateAPIView):
    serializer_class = DamageReportSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    pagination_class = None

    def get_queryset(self):
        queryset = DamageReport.objects.order_by('-calamity_date')
        barangay = self.request.query_params.get('barangay')
        season = self.request.query_params.get('season')
        year = self.request.query_params.get('year')
        if barangay:
            queryset = queryset.filter(barangay__iexact=barangay)
        if season:
            queryset = queryset.filter(season__iexact=season)
        if year:
            queryset = queryset.filter(year=year)
        return queryset

    def perform_create(self, serializer):
        serializer.save(recorded_by=self.request.user)


class DamageReportDetailView(generics.RetrieveUpdateDestroyAPIView):
    queryset = DamageReport.objects.all()
    serializer_class = DamageReportSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    http_method_names = ['get', 'patch', 'delete']
    lookup_field = 'pk'


class YearlyProductionTrendView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def get(self, request):
        current_year = timezone.now().year
        start_year = current_year - 4

        records = (
            HarvestRecord.objects
            .filter(year__gte=start_year, year__lte=current_year)
            .values('year')
            .annotate(total_volume_kg=Sum('volume_kg'))
            .order_by('year')
        )

        data_by_year = {item['year']: item['total_volume_kg'] for item in records}

        return Response([
            {
                'year': year,
                'total_volume_kg': data_by_year.get(year, 0),
            }
            for year in range(start_year, current_year + 1)
        ])


class DashboardExtendedSummaryView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def get(self, request):
        # Fisherfolk counts
        fisherfolk_active = Fisherfolk.objects.filter(is_active=True).count()
        fisherfolk_breakdown = {
            'aquaculture': Fisherfolk.objects.filter(fisherfolk_type='Aquaculture').count(),
            'capture': Fisherfolk.objects.filter(fisherfolk_type='Capture').count(),
            'vendor': Fisherfolk.objects.filter(fisherfolk_type='Vendor').count(),
        }

        return Response({
            'fisherfolk_active': fisherfolk_active,
            'fisherfolk_breakdown': fisherfolk_breakdown,
        })


# ─── Module 1: CropType ──────────────────────────────────────────────────────

class CropTypeListView(generics.ListCreateAPIView):
    serializer_class = CropTypeSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    pagination_class = None
    queryset = CropType.objects.all().order_by('name')


class CropTypeDetailView(generics.RetrieveUpdateDestroyAPIView):
    serializer_class = CropTypeSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    queryset = CropType.objects.all()


# ─── Module 2: CropMonitoringRecord ──────────────────────────────────────────

class CropMonitoringListCreateView(generics.ListCreateAPIView):
    serializer_class = CropMonitoringRecordSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def get_queryset(self):
        qs = CropMonitoringRecord.objects.select_related(
            'farm_parcel', 'farm_parcel__farmer', 'farmer', 'crop_type', 'recorded_by', 'farmers_association',
        ).prefetch_related('activity_logs', 'labor_logs', 'harvest_records').order_by('-date_recorded')
        qs = qs.annotate(
            has_disaster=Exists(
                CropDisasterReport.objects.filter(
                    monitoring_record=OuterRef('pk'), status='Active',
                )
            )
        )
        params = self.request.query_params
        if params.get('year'):
            qs = qs.filter(year=params['year'])
        if params.get('season'):
            qs = qs.filter(season__iexact=params['season'])
        if params.get('barangay'):
            qs = qs.filter(farm_parcel__barangay=params['barangay'])
        if params.get('status'):
            qs = qs.filter(status=params['status'])
        if params.get('farm_parcel'):
            qs = qs.filter(farm_parcel_id=params['farm_parcel'])
        return qs

    def perform_create(self, serializer):
        serializer.save(recorded_by=self.request.user)


class CropMonitoringDetailView(generics.RetrieveUpdateDestroyAPIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    queryset = CropMonitoringRecord.objects.select_related(
        'farm_parcel', 'farm_parcel__farmer', 'crop_type',
    ).prefetch_related('activity_logs', 'input_logs', 'pest_disease_reports', 'labor_logs')

    def get_serializer_class(self):
        if self.request.method == 'GET':
            return CropMonitoringDetailSerializer
        return CropMonitoringRecordSerializer

    def destroy(self, request, *args, **kwargs):
        if request.user.role != 'ADMIN':
            return Response(
                {'detail': 'Only administrators can delete crop records.'},
                status=drf_status.HTTP_403_FORBIDDEN,
            )
        return super().destroy(request, *args, **kwargs)


# ─── Crop Duplicate Check ─────────────────────────────────────────────────────

class CropDuplicateCheckView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def get(self, request):
        farmer_id   = request.query_params.get('farmer')
        crop_type_id = request.query_params.get('crop_type')
        if not farmer_id or not crop_type_id:
            return Response({'duplicates': []})
        qs = CropMonitoringRecord.objects.filter(
            farmer_id=farmer_id,
            crop_type_id=crop_type_id,
            status=CropMonitoringRecord.Status.ACTIVE,
        ).select_related('crop_type').order_by('-date_planted')[:5]
        results = [
            {
                'id':           r.id,
                'crop_type':    r.crop_type.name if r.crop_type else '',
                'date_planted': str(r.date_planted) if r.date_planted else None,
                'area_planted_ha': str(r.area_planted_ha),
                'status':       r.status,
            }
            for r in qs
        ]
        return Response({'duplicates': results})


# ─── Module 2b: Crop Activity Suggestions ────────────────────────────────────

class CropActivitySuggestionsView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def get(self, request, monitoring_id):
        from datetime import date
        from django.shortcuts import get_object_or_404
        record = get_object_or_404(
            CropMonitoringRecord.objects.select_related('crop_type'),
            pk=monitoring_id,
        )

        from datetime import timedelta

        stages = CropActivityStage.objects.filter(
            crop_type=record.crop_type
        ).order_by('day_after_planting')

        activity_logs = CropActivityLog.objects.filter(monitoring_record=record).values(
            'activity_type', 'activity_date', 'target_day'
        )

        today = date.today()
        days_elapsed = (today - record.date_planted).days if record.date_planted else 0
        total_days = record.crop_type.days_to_harvest if record.crop_type else 0
        planted = record.date_planted

        suggestions = []
        for stage in stages:
            days_list = stage.application_days if stage.application_days else [stage.day_after_planting]
            for occurrence_day in days_list:
                due_date = planted + timedelta(days=occurrence_day) if planted else None
                days_diff = occurrence_day - days_elapsed

                is_done = False
                done_date = None
                if due_date:
                    window = stage.window_days
                    for log in activity_logs:
                        if log['activity_type'] != stage.stage_name:
                            continue
                        target_match = log.get('target_day') is not None and log['target_day'] == occurrence_day
                        date_match = abs((log['activity_date'] - due_date).days) <= window
                        if target_match or date_match:
                            is_done = True
                            done_date = log['activity_date'].isoformat()
                            break

                if is_done:
                    status = 'done'
                elif days_diff < -stage.window_days:
                    status = 'overdue'
                elif abs(days_diff) <= stage.window_days:
                    status = 'due_now'
                else:
                    status = 'upcoming'

                suggestions.append({
                    'id': stage.id,
                    'occurrence_key': f'{stage.id}_{occurrence_day}',
                    'activity_name': stage.stage_name,
                    'day_after_planting': occurrence_day,
                    'recommendation': stage.recommendation,
                    'priority': stage.priority,
                    'window_days': stage.window_days,
                    'status': status,
                    'is_completed': is_done,
                    'done_date': done_date,
                    'days_diff': days_diff,
                })

        suggestions.sort(key=lambda s: s['day_after_planting'])

        # Disaster response rules — only ACTIVE disasters
        active_disasters = list(
            CropDisasterReport.objects.filter(
                monitoring_record=record, status='Active',
            ).order_by('-date_of_occurrence')
        )
        disaster_response = []
        for disaster in active_disasters:
            rules_qs = DisasterResponseRule.objects.filter(
                disaster_type=disaster.disaster_type,
            ).filter(
                Q(crop_type=record.crop_type) | Q(crop_type__isnull=True)
            ).order_by('-priority')

            rules_list = []
            for rule in rules_qs:
                is_done = any(
                    log['activity_type'] == rule.title and
                    log['activity_date'] >= disaster.date_of_occurrence
                    for log in activity_logs
                )
                rules_list.append({
                    'disaster_type': rule.disaster_type,
                    'title': rule.title,
                    'action': rule.action,
                    'priority': rule.priority,
                    'is_done': is_done,
                })

            all_rules_done = len(rules_list) > 0 and all(r['is_done'] for r in rules_list)

            disaster_response.append({
                'disaster_id': disaster.id,
                'disaster_type': disaster.disaster_type,
                'date_of_occurrence': disaster.date_of_occurrence.isoformat(),
                'status': disaster.status,
                'all_rules_done': all_rules_done,
                'rules': rules_list,
            })

        return Response({
            'days_elapsed': days_elapsed,
            'total_days': total_days,
            'suggestions': suggestions,
            'disaster_response': disaster_response,
        })


# ─── Module 3: CropActivityLog ───────────────────────────────────────────────

class CropActivityLogListCreateView(generics.ListCreateAPIView):
    serializer_class = CropActivityLogSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    pagination_class = None

    def get_queryset(self):
        monitoring_id = self.kwargs.get('monitoring_id')
        return CropActivityLog.objects.filter(
            monitoring_record_id=monitoring_id,
        ).order_by('-activity_date')

    def perform_create(self, serializer):
        monitoring_id = self.kwargs.get('monitoring_id')
        serializer.save(
            monitoring_record_id=monitoring_id,
            recorded_by=self.request.user,
        )


class CropActivityLogDeleteView(generics.DestroyAPIView):
    serializer_class = CropActivityLogSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    queryset = CropActivityLog.objects.all()


# ─── CropDisasterReport ──────────────────────────────────────────────────────

class CropDisasterReportListCreateView(generics.ListCreateAPIView):
    serializer_class = CropDisasterReportSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    pagination_class = None

    def get_queryset(self):
        monitoring_id = self.kwargs.get('monitoring_id')
        return CropDisasterReport.objects.filter(
            monitoring_record_id=monitoring_id,
        ).order_by('-date_of_occurrence')

    def perform_create(self, serializer):
        monitoring_id = self.kwargs.get('monitoring_id')
        serializer.save(
            monitoring_record_id=monitoring_id,
            reported_by=self.request.user,
        )


class CropDisasterResolveView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def post(self, request, pk):
        from datetime import date as _date
        disaster = get_object_or_404(CropDisasterReport, pk=pk)
        if disaster.status == 'Resolved':
            return Response({'detail': 'Disaster is already resolved.'}, status=400)

        record = disaster.monitoring_record
        rules_qs = DisasterResponseRule.objects.filter(
            disaster_type=disaster.disaster_type,
        ).filter(
            Q(crop_type=record.crop_type) | Q(crop_type__isnull=True)
        )
        activity_logs = CropActivityLog.objects.filter(
            monitoring_record=record
        ).values('activity_type', 'activity_date')

        all_done = True
        has_rules = False
        for rule in rules_qs:
            has_rules = True
            is_done = any(
                log['activity_type'] == rule.title and
                log['activity_date'] >= disaster.date_of_occurrence
                for log in activity_logs
            )
            if not is_done:
                all_done = False
                break

        if has_rules and not all_done:
            return Response(
                {'detail': 'Cannot resolve: not all recommended response rules have been completed.'},
                status=400,
            )

        disaster.status = 'Resolved'
        disaster.resolved_date = _date.today()
        disaster.resolved_by = request.user
        disaster.save(update_fields=['status', 'resolved_date', 'resolved_by'])
        serializer = CropDisasterReportSerializer(disaster)
        return Response(serializer.data, status=200)


# ─── LaborOperationLog ──────────────────────────────────────────────────────

class LaborOperationLogListCreateView(generics.ListCreateAPIView):
    serializer_class = LaborOperationLogSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    pagination_class = None

    def get_queryset(self):
        monitoring_id = self.kwargs.get('monitoring_id')
        return LaborOperationLog.objects.filter(
            monitoring_record_id=monitoring_id,
        ).order_by('-date')

    def perform_create(self, serializer):
        monitoring_id = self.kwargs.get('monitoring_id')
        serializer.save(
            monitoring_record_id=monitoring_id,
            logged_by=self.request.user,
        )


class LaborOperationLogDeleteView(generics.DestroyAPIView):
    serializer_class = LaborOperationLogSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    queryset = LaborOperationLog.objects.all()


# ─── FarmInputLog ───────────────────────────────────────────────────────────

class FarmInputLogListCreateView(generics.ListCreateAPIView):
    serializer_class = FarmInputLogSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    pagination_class = None

    def get_queryset(self):
        monitoring_id = self.kwargs.get('monitoring_id')
        return FarmInputLog.objects.filter(
            monitoring_record_id=monitoring_id,
        ).order_by('-application_date')

    def perform_create(self, serializer):
        serializer.save(
            monitoring_record_id=self.kwargs.get('monitoring_id'),
            recorded_by=self.request.user,
        )


class FarmInputLogDeleteView(generics.DestroyAPIView):
    serializer_class = FarmInputLogSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    queryset = FarmInputLog.objects.all()


# ─── PestDiseaseReport ───────────────────────────────────────────────────────

class PestDiseaseReportListCreateView(generics.ListCreateAPIView):
    serializer_class = PestDiseaseReportSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    pagination_class = None

    def get_queryset(self):
        monitoring_id = self.kwargs.get('monitoring_id')
        return PestDiseaseReport.objects.filter(
            monitoring_record_id=monitoring_id,
        ).order_by('-report_date')

    def perform_create(self, serializer):
        serializer.save(
            monitoring_record_id=self.kwargs.get('monitoring_id'),
            recorded_by=self.request.user,
        )


class PestDiseaseReportDeleteView(generics.DestroyAPIView):
    serializer_class = PestDiseaseReportSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    queryset = PestDiseaseReport.objects.all()


# ─── Module 5: Monitoring Analytics ──────────────────────────────────────────

class SeasonalVarianceView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def get(self, request):
        year = request.query_params.get('year')
        qs = CropMonitoringRecord.objects.filter(
            status=CropMonitoringRecord.Status.HARVESTED,
            actual_yield_kg__isnull=False,
            estimated_yield_kg__isnull=False,
        )
        if year:
            qs = qs.filter(year=year)
        data = (
            qs.values('season', 'crop_type__name')
            .annotate(
                total_estimated=Sum('estimated_yield_kg'),
                total_actual=Sum('actual_yield_kg'),
            )
            .order_by('season', 'crop_type__name')
        )
        return Response(list(data))


class YearOnYearVarianceView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def get(self, request):
        crop_type = request.query_params.get('crop_type')
        qs = CropMonitoringRecord.objects.filter(
            status=CropMonitoringRecord.Status.HARVESTED,
            actual_yield_kg__isnull=False,
            estimated_yield_kg__isnull=False,
        )
        if crop_type:
            qs = qs.filter(crop_type__name__icontains=crop_type)
        data = (
            qs.values('year', 'crop_type__name')
            .annotate(
                total_estimated=Sum('estimated_yield_kg'),
                total_actual=Sum('actual_yield_kg'),
            )
            .order_by('year', 'crop_type__name')
        )
        return Response(list(data))


# ─── Module 2: BarangayCropSummary ───────────────────────────────────────────

class BarangayCropSummaryListCreateView(generics.ListCreateAPIView):
    serializer_class = BarangayCropSummarySerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    pagination_class = None

    def get_queryset(self):
        qs = BarangayCropSummary.objects.select_related(
            'recorded_by', 'causal_damage_report',
        ).order_by('barangay', 'crop_name')
        p = self.request.query_params
        if p.get('barangay'):
            qs = qs.filter(barangay=p['barangay'])
        if p.get('crop_name'):
            qs = qs.filter(crop_name=p['crop_name'])
        if p.get('season'):
            qs = qs.filter(season=p['season'])
        if p.get('year'):
            qs = qs.filter(year=p['year'])
        if p.get('alert_level'):
            qs = qs.filter(alert_level=p['alert_level'])
        if p.get('submission_status'):
            qs = qs.filter(submission_status=p['submission_status'])
        if p.get('stage'):
            qs = qs.filter(current_growth_stage=p['stage'])
        return qs

    def perform_create(self, serializer):
        serializer.save(recorded_by=self.request.user)


class BarangayCropSummaryDetailView(generics.RetrieveUpdateDestroyAPIView):
    serializer_class = BarangayCropSummarySerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    queryset = BarangayCropSummary.objects.select_related('recorded_by', 'causal_damage_report')
    http_method_names = ['get', 'patch', 'delete']


class BarangayCropSummaryLinkRecordsView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def post(self, request, pk):
        summary = BarangayCropSummary.objects.get(pk=pk)
        ids = request.data.get('monitoring_record_ids', [])
        records = CropMonitoringRecord.objects.filter(pk__in=ids)
        summary.linked_monitoring_records.add(*records)
        serializer = BarangayCropSummarySerializer(summary)
        return Response(serializer.data)


ALERT_ORDER = {'None': 0, 'Low': 1, 'Medium': 2, 'High': 3}
STATUS_ORDER = {'Draft': 0, 'Submitted': 1, 'Verified': 2}


class MunicipalityCropDashboardView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def get(self, request):
        import datetime
        today = datetime.date.today()
        month = today.month
        default_season = 'Wet' if 6 <= month <= 11 else 'Dry'
        season = request.query_params.get('season', default_season)
        year = int(request.query_params.get('year', today.year))

        qs = BarangayCropSummary.objects.filter(season=season, year=year)

        # Overall totals
        agg = qs.aggregate(
            total_area=Sum('area_planted_ha'),
            total_estimated=Sum('estimated_yield_kg'),
            total_actual=Sum('actual_yield_kg'),
        )
        total_area = agg['total_area'] or Decimal('0')
        total_est = agg['total_estimated'] or Decimal('0')
        total_act = agg['total_actual']

        overall_variance_pct = None
        if total_act is not None and total_est and total_est != 0:
            overall_variance_pct = float(((total_act - total_est) / total_est) * 100)

        # Stage distribution per crop
        stage_qs = qs.values('crop_name', 'current_growth_stage').annotate(cnt=Count('id'))
        stage_by_crop = defaultdict(dict)
        for row in stage_qs:
            if row['current_growth_stage']:
                stage_by_crop[row['crop_name']][row['current_growth_stage']] = row['cnt']

        # Alert count per crop
        alert_qs = qs.filter(alert_level__in=['Medium', 'High']).values('crop_name').annotate(
            cnt=Count('id')
        )
        alert_by_crop = {row['crop_name']: row['cnt'] for row in alert_qs}

        # By crop
        crop_qs = qs.values('crop_name').annotate(
            total_area_ha=Sum('area_planted_ha'),
            barangay_count=Count('barangay', distinct=True),
            estimated_yield_kg=Sum('estimated_yield_kg'),
            actual_yield_kg=Sum('actual_yield_kg'),
        ).order_by('-total_area_ha')

        by_crop = []
        for row in crop_qs:
            cn = row['crop_name']
            est = row['estimated_yield_kg'] or Decimal('0')
            act = row['actual_yield_kg']
            var_pct = None
            if act is not None and est and est != 0:
                var_pct = float(((act - est) / est) * 100)
            by_crop.append({
                'crop_name': cn,
                'total_area_ha': float(row['total_area_ha'] or 0),
                'barangay_count': row['barangay_count'],
                'estimated_yield_kg': float(est),
                'actual_yield_kg': float(act) if act is not None else None,
                'variance_pct': var_pct,
                'alert_count': alert_by_crop.get(cn, 0),
                'stage_distribution': stage_by_crop.get(cn, {}),
            })

        # By barangay — Python-level to get max alert_level
        all_records = list(qs.values(
            'barangay', 'alert_level', 'submission_status',
            'area_planted_ha', 'crop_name', 'observed_problems',
        ))
        barangay_map = {}
        for rec in all_records:
            b = rec['barangay']
            if b not in barangay_map:
                barangay_map[b] = {
                    'barangay': b,
                    'crops_reported': 0,
                    'total_area_ha': Decimal('0'),
                    'alert_level': 'None',
                    'submission_status': 'Draft',
                    'has_observed_problems': False,
                }
            d = barangay_map[b]
            d['crops_reported'] += 1
            d['total_area_ha'] += (rec['area_planted_ha'] or Decimal('0'))
            if ALERT_ORDER.get(rec['alert_level'], 0) > ALERT_ORDER.get(d['alert_level'], 0):
                d['alert_level'] = rec['alert_level']
            if STATUS_ORDER.get(rec['submission_status'], 0) > STATUS_ORDER.get(d['submission_status'], 0):
                d['submission_status'] = rec['submission_status']
            if rec['observed_problems']:
                d['has_observed_problems'] = True

        for d in barangay_map.values():
            d['total_area_ha'] = float(d['total_area_ha'])

        by_barangay = sorted(barangay_map.values(), key=lambda x: x['barangay'])

        # Submission tracker
        all_barangay_names = [b[0] for b in BARANGAY_CHOICES]
        reported = set(barangay_map.keys())
        not_submitted_list = [b for b in all_barangay_names if b not in reported]

        submission_tracker = {
            'total_barangays': len(all_barangay_names),
            'submitted': len(reported),
            'not_submitted': len(not_submitted_list),
            'not_submitted_list': not_submitted_list,
        }

        # Alerts list
        alerts = list(
            qs.filter(alert_level__in=['Medium', 'High'])
            .values('barangay', 'crop_name', 'alert_level', 'observed_problems')
            .order_by('-alert_level')
        )

        return Response({
            'current_season': season,
            'current_year': year,
            'total_area_planted_ha': float(total_area),
            'total_estimated_yield_kg': float(total_est),
            'total_actual_yield_kg': float(total_act) if total_act is not None else None,
            'overall_variance_pct': overall_variance_pct,
            'by_crop': by_crop,
            'by_barangay': by_barangay,
            'submission_tracker': submission_tracker,
            'alerts': alerts,
        })


class CropHistoricalTimelineView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def get(self, request):
        crop_name = request.query_params.get('crop_name')
        barangay = request.query_params.get('barangay')

        if not crop_name:
            return Response({'detail': 'crop_name is required.'}, status=400)

        qs = BarangayCropSummary.objects.filter(crop_name=crop_name)
        if barangay:
            qs = qs.filter(barangay=barangay)

        # Group by season + year
        grouped = qs.values('season', 'year').annotate(
            total_area_ha=Sum('area_planted_ha'),
            total_estimated_yield_kg=Sum('estimated_yield_kg'),
            total_actual_yield_kg=Sum('actual_yield_kg'),
            barangay_count=Count('barangay', distinct=True),
        ).order_by('year', 'season')

        # Primary causal factor per season+year (most common)
        causal_qs = qs.exclude(causal_factor='').exclude(causal_factor='None').values(
            'season', 'year', 'causal_factor'
        ).annotate(cnt=Count('id')).order_by('year', 'season', '-cnt')

        causal_map = {}
        for row in causal_qs:
            key = (row['season'], row['year'])
            if key not in causal_map:
                causal_map[key] = row['causal_factor']

        timeline = []
        for row in grouped:
            key = (row['season'], row['year'])
            est = row['total_estimated_yield_kg']
            act = row['total_actual_yield_kg']
            var_pct = None
            if act is not None and est and est != 0:
                var_pct = float(((Decimal(str(act)) - Decimal(str(est))) / Decimal(str(est))) * 100)
            timeline.append({
                'season': row['season'],
                'year': row['year'],
                'total_area_ha': float(row['total_area_ha'] or 0),
                'total_estimated_yield_kg': float(est) if est is not None else None,
                'total_actual_yield_kg': float(act) if act is not None else None,
                'variance_pct': var_pct,
                'primary_causal_factor': causal_map.get(key, 'None'),
                'barangay_count': row['barangay_count'],
            })

        return Response({
            'crop_name': crop_name,
            'scope': barangay if barangay else 'Municipality-wide',
            'timeline': timeline,
        })


# ─── Module 3: Submission Status Tracker ─────────────────────────────────────

class BarangaySubmissionStatusView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def get(self, request):
        season = request.query_params.get('season')
        year = request.query_params.get('year')
        if not season or not year:
            return Response({'detail': 'season and year are required.'}, status=400)

        all_barangays = [b[0] for b in BARANGAY_CHOICES]
        records = BarangayCropSummary.objects.filter(
            season=season, year=year,
        ).values('barangay', 'crop_name', 'submission_status')

        by_barangay = defaultdict(list)
        for rec in records:
            by_barangay[rec['barangay']].append(rec)

        result = []
        submitted = 0
        not_submitted = 0
        verified = 0

        for b in all_barangays:
            if b in by_barangay:
                recs = by_barangay[b]
                crops = [r['crop_name'] for r in recs]
                statuses = [r['submission_status'] for r in recs]
                if all(s == 'Verified' for s in statuses):
                    status = 'Verified'
                    verified += 1
                elif any(s in ('Submitted', 'Verified') for s in statuses):
                    status = 'Submitted'
                else:
                    status = 'Draft'
                submitted += 1
                result.append({
                    'barangay': b,
                    'status': status,
                    'crops_reported': crops,
                    'record_count': len(recs),
                })
            else:
                not_submitted += 1
                result.append({
                    'barangay': b,
                    'status': 'Not Submitted',
                    'crops_reported': [],
                    'record_count': 0,
                })

        return Response({
            'season': season,
            'year': int(year),
            'summary': {
                'submitted': submitted,
                'not_submitted': not_submitted,
                'verified': verified,
            },
            'barangays': result,
        })


# ─── Crop Category / Subcategory ─────────────────────────────────────────────

class CropCategoryListCreateView(generics.ListCreateAPIView):
    serializer_class = CropCategorySerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    pagination_class = None
    queryset = CropCategory.objects.order_by('name')

    def perform_create(self, serializer):
        serializer.save(created_by=self.request.user)


class CropCategoryDetailView(generics.RetrieveUpdateDestroyAPIView):
    serializer_class = CropCategorySerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    queryset = CropCategory.objects.all()


class CropSubcategoryListCreateView(generics.ListCreateAPIView):
    serializer_class = CropSubcategorySerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    pagination_class = None

    def get_queryset(self):
        qs = CropSubcategory.objects.select_related('category').order_by('name')
        category_id = self.request.query_params.get('category_id')
        if category_id:
            qs = qs.filter(category_id=category_id)
        return qs

    def perform_create(self, serializer):
        serializer.save(created_by=self.request.user)


class CropSubcategoryDetailView(generics.RetrieveUpdateDestroyAPIView):
    serializer_class = CropSubcategorySerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    queryset = CropSubcategory.objects.select_related('category').all()


# ─── CropRemark ──────────────────────────────────────────────────────────────

class CropRemarkListCreateView(generics.ListCreateAPIView):
    serializer_class = CropRemarkSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    pagination_class = None

    def get_queryset(self):
        qs = CropRemark.objects.select_related('crop_type', 'recorded_by').order_by('-date_recorded')
        crop_type_id = self.request.query_params.get('crop_type')
        if crop_type_id:
            qs = qs.filter(crop_type_id=crop_type_id)
        return qs

    def perform_create(self, serializer):
        serializer.save(recorded_by=self.request.user)


class CropRemarkDetailView(generics.RetrieveUpdateDestroyAPIView):
    serializer_class = CropRemarkSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    queryset = CropRemark.objects.select_related('crop_type', 'recorded_by')
    http_method_names = ['get', 'patch', 'delete']


class MonitoringRemarkListCreateView(generics.ListCreateAPIView):
    serializer_class = MonitoringRemarkSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    pagination_class = None

    def get_queryset(self):
        return CropRemark.objects.filter(
            monitoring_record_id=self.kwargs['monitoring_id']
        ).select_related('recorded_by').order_by('-date_recorded')

    def perform_create(self, serializer):
        from django.shortcuts import get_object_or_404
        monitoring_record = get_object_or_404(CropMonitoringRecord, pk=self.kwargs['monitoring_id'])
        serializer.save(
            recorded_by=self.request.user,
            monitoring_record=monitoring_record,
            crop_type=monitoring_record.crop_type,
        )


# ─── CropActivityStage CRUD ──────────────────────────────────────────────────

class CropActivityStageListCreateView(generics.ListCreateAPIView):
    serializer_class = CropActivityStageSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    pagination_class = None

    def get_queryset(self):
        return CropActivityStage.objects.filter(
            crop_type_id=self.kwargs['crop_type_id']
        ).order_by('day_after_planting')

    def perform_create(self, serializer):
        serializer.save(crop_type_id=self.kwargs['crop_type_id'])


class CropActivityStageDeleteView(generics.DestroyAPIView):
    serializer_class = CropActivityStageSerializer
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    queryset = CropActivityStage.objects.all()


# ─── Forecasting Template & Import ───────────────────────────────────────────

class ForecastingTemplateView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]

    def get(self, request):
        from django.http import HttpResponse
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = 'Forecasting Import'

        headers = ['Association Name', 'Crop Name', 'Date Planted (YYYY-MM-DD)', 'Area Planted (ha)']
        green_fill = PatternFill(start_color='1E4620', end_color='1E4620', fill_type='solid')
        white_font = Font(color='FFFFFF', bold=True)

        col_widths = [35, 25, 28, 20]
        for col, (header, width) in enumerate(zip(headers, col_widths), 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = green_fill
            cell.font = white_font
            ws.column_dimensions[cell.column_letter].width = width

        ws.cell(row=2, column=1, value='Aringay Farmers Association')
        ws.cell(row=2, column=2, value='Palay (Rice)')
        ws.cell(row=2, column=3, value='2025-01-15')
        ws.cell(row=2, column=4, value=1.50)

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)

        response = HttpResponse(
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = 'attachment; filename="forecasting_template.xlsx"'
        return response


class ForecastingImportView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrEvaluator]
    parser_classes = [MultiPartParser]

    def post(self, request):
        import datetime as dt

        from openpyxl import load_workbook

        from associations.models import FarmersAssociation

        file = request.FILES.get('file')
        if not file:
            return Response({'error': 'No file uploaded.'}, status=400)

        try:
            wb = load_workbook(file, data_only=True)
            ws = wb.active
        except Exception as exc:
            return Response({'error': f'Invalid Excel file: {exc}'}, status=400)

        created = 0
        errors = []

        for row_num, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
            if not any(cell for cell in row if cell is not None):
                continue

            assoc_name = str(row[0]).strip() if row[0] is not None else ''
            crop_name  = str(row[1]).strip() if row[1] is not None else ''
            date_str   = str(row[2]).strip() if row[2] is not None else ''
            area_str   = str(row[3]).strip() if row[3] is not None else ''

            if not assoc_name or not crop_name or not date_str:
                errors.append({'row': row_num, 'error': 'Association Name, Crop Name, and Date Planted are required.'})
                continue

            try:
                assoc = FarmersAssociation.objects.filter(name__iexact=assoc_name).first()
                if not assoc:
                    errors.append({'row': row_num, 'error': f'Association "{assoc_name}" not found.'})
                    continue
            except Exception as exc:
                errors.append({'row': row_num, 'error': str(exc)})
                continue

            try:
                crop_type = CropType.objects.filter(name__iexact=crop_name).first()
                if not crop_type:
                    errors.append({'row': row_num, 'error': f'Crop type "{crop_name}" not found.'})
                    continue
            except Exception as exc:
                errors.append({'row': row_num, 'error': str(exc)})
                continue

            try:
                date_planted = dt.date.fromisoformat(date_str[:10])
            except ValueError:
                errors.append({'row': row_num, 'error': f'Invalid date "{date_str}". Use YYYY-MM-DD.'})
                continue

            try:
                area = float(area_str) if area_str else None
            except ValueError:
                errors.append({'row': row_num, 'error': f'Invalid area value "{area_str}".'})
                continue

            try:
                month = date_planted.month
                season = 'Wet' if 5 <= month <= 10 else 'Dry'
                CropMonitoringRecord.objects.create(
                    farmers_association=assoc,
                    crop_type=crop_type,
                    date_planted=date_planted,
                    area_planted_ha=area,
                    season=season,
                    year=date_planted.year,
                    recorded_by=request.user,
                )
                created += 1
            except Exception as exc:
                errors.append({'row': row_num, 'error': str(exc)})

        return Response({'created': created, 'errors': errors})
